"""Revisão OOXML do REFAT: preserva fórmulas e seus resultados armazenados."""
from collections import OrderedDict
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import posixpath
import re
from tempfile import NamedTemporaryFile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape, quoteattr
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from .parsers.financeiro_geral import parse_financeiro_geral
from .project_status import STATUS_PROJETOS, etapa_do_status, farol_financeiro


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
TABLES = (("tbl_ProjetosPDI", "tbl_VisaoPDI"), ("tbl_ProjetosFCRH", "tbl_VisaoFCRH"), ("tbl_ProjetosACS", "tbl_VisaoACS"))
CELL = re.compile(r'<c\b[^>]*\br="([A-Z]+)(\d+)"[^>]*?(?:/>|>.*?</c>)', re.S)
ROW = re.compile(r'(<row\b[^>]*\br="(\d+)"[^>]*>)(.*?)(</row>)', re.S)


def _texto_cell(ref, texto, style=None, formula=None):
    estilo = f' s="{style}"' if style is not None else ""
    if formula is not None:
        return f'<c r="{ref}"{estilo} t="str"><f>{escape(formula)}</f><v>{escape(texto)}</v></c>'
    return f'<c r="{ref}"{estilo} t="inlineStr"><is><t xml:space="preserve">{escape(texto)}</t></is></c>'


def _mover_refs(formula, inicio, fim, deslocamento):
    def trocar(match):
        coluna = column_index_from_string(match[2])
        if inicio <= coluna <= fim and int(match[4]) <= 2:
            return match[1] + get_column_letter(coluna + deslocamento) + match[3] + match[4]
        return match[0]
    return re.sub(r'(\$?)([A-Z]{1,3})(\$?)([1-9][0-9]*)\b', trocar, formula)


def _ref_intervalo(ref, deslocamento):
    a, r1, b, r2 = range_boundaries(ref)
    return f"{get_column_letter(a+deslocamento)}{r1}:{get_column_letter(b+deslocamento)}{r2}"


def revisar_planilha(caminho, backup_dir=None):
    caminho = Path(caminho)
    original = caminho.read_bytes()
    assinatura = sha256(original).hexdigest()
    antes = parse_financeiro_geral(BytesIO(original))
    if antes["erros"]:
        raise ValueError("A planilha original não passou na validação: " + " | ".join(antes["erros"]))
    with ZipFile(caminho) as arquivo:
        infos = arquivo.infolist()
        partes = OrderedDict((i.filename, arquivo.read(i.filename)) for i in infos)
    tabelas = {ET.fromstring(v).attrib["displayName"]: nome for nome, v in partes.items() if nome.startswith("xl/tables/") and nome.endswith(".xml")}
    if all("ETAPA DO PROJETO" in partes[tabelas[nome]].decode("utf-8") for nome, _ in TABLES):
        return {"alterado": False, "sha256": assinatura, "projetos": len(antes["projetos"])}
    if backup_dir is not None:
        backup_dir = Path(backup_dir)
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup = backup_dir / f"refat-v2-antes-{assinatura[:16]}.xlsx"
        if not backup.exists():
            backup.write_bytes(original)

    wb = load_workbook(BytesIO(original), data_only=True)
    book = ET.fromstring(partes["xl/workbook.xml"])
    relationships = ET.fromstring(partes["xl/_rels/workbook.xml.rels"])
    destinos = {r.attrib["Id"]: posixpath.normpath("xl/" + r.attrib["Target"]) if not r.attrib["Target"].startswith("/") else r.attrib["Target"].lstrip("/") for r in relationships}
    sheets = {s.attrib["name"]: destinos[s.attrib[f"{{{REL}}}id"]] for s in book.find(f"{{{NS}}}sheets")}
    statuses = list(STATUS_PROJETOS)
    report = []
    for nome, nome_resumo in TABLES:
        ws = next(s for s in wb if nome in s.tables)
        tabela = ws.tables[nome]
        a, inicio, ultima, fim = range_boundaries(tabela.ref)
        headers = [c.name for c in tabela.tableColumns]
        indice_status = headers.index("STATUS") + a
        inicio_col = next((i+a for i, h in enumerate(headers) if h.startswith("INICIO")), None)
        final_col = next(i+a for i, h in enumerate(headers) if h in ("FINAL", "FINAL DO PROJETO", "DATA FINAL DO PROJETO"))
        orcado_col = next(i+a for i, h in enumerate(headers) if "ORÇADO" in h and "DIFERENÇA" not in h and not h.startswith("%"))
        realizado_col = next(i+a for i, h in enumerate(headers) if h.startswith("REALIZADO") and "+" not in h)
        novos = ([] if inicio_col else ["ASSINATURA DA SÍNTESE"]) + ["DATA DO KICK-OFF", "FINAL PREVISTO REAL", "FAROL", "ETAPA DO PROJETO"]
        nova_ultima = ultima + len(novos)
        resumo = ws.tables[nome_resumo]
        resumo_inicio, _, resumo_fim, _ = range_boundaries(resumo.ref)
        deslocamento = max(0, nova_ultima - resumo_inicio + 1)
        worksheet_path = sheets[ws.title]
        xml = partes[worksheet_path].decode("utf-8")
        dados = {r: [ws.cell(r, c).value for c in range(a, ultima+1)] for r in range(inicio+1, fim+1)}
        header_style = ws.cell(inicio, final_col).style_id
        date_style = ws.cell(inicio+1, final_col).style_id
        origem_grupos = []
        formulas_novas = {}

        def revisar_linha(match):
            numero = int(match[2])
            cells = {}
            residual = CELL.sub("", match[3]).strip()
            if residual:
                raise ValueError("Conteúdo de linha desconhecido; a revisão foi interrompida sem gravar.")
            for cell in CELL.finditer(match[3]):
                coluna = column_index_from_string(cell[1])
                valor = cell[0]
                if deslocamento and numero <= 2 and resumo_inicio <= coluna <= resumo_fim:
                    coluna += deslocamento
                    valor = re.sub(r'\br="[A-Z]+[0-9]+"', f'r="{get_column_letter(coluna)}{numero}"', valor, count=1)
                    valor = re.sub(r'(<f\b[^>]*>)(.*?)(</f>)', lambda f: f[1]+_mover_refs(f[2], resumo_inicio, resumo_fim, deslocamento)+f[3], valor, flags=re.S)
                cells[coluna] = valor
            if numero == inicio:
                for coluna, label in ((inicio_col, "ASSINATURA DA SÍNTESE"), (final_col, "FINAL PREVISTO NA SÍNTESE")):
                    if coluna:
                        cells[coluna] = _texto_cell(f"{get_column_letter(coluna)}{numero}", label, header_style)
            if inicio <= numero <= fim:
                for indice, label in enumerate(novos, start=ultima+1):
                    ref = f"{get_column_letter(indice)}{numero}"
                    if numero == inicio:
                        cells[indice] = _texto_cell(ref, label, header_style)
                    elif label in ("ASSINATURA DA SÍNTESE", "DATA DO KICK-OFF", "FINAL PREVISTO REAL"):
                        cells[indice] = f'<c r="{ref}" s="{date_style}"/>'
                    else:
                        values = dados[numero]
                        status = values[indice_status-a] or ""
                        if status and status not in statuses:
                            statuses.append(status)
                        if label == "FAROL":
                            orcado, realizado = values[orcado_col-a], values[realizado_col-a]
                            value = farol_financeiro(realizado, orcado)
                            o, r = f"{get_column_letter(orcado_col)}{numero}", f"{get_column_letter(realizado_col)}{numero}"
                            formula = f'IF(OR(ISBLANK({r}),ISBLANK({o}),{o}<=0),"Sem lançamento",IF({r}/{o}>=0.9,"Meta atingida",IF({r}/{o}>=0.5,"Execução parcial","Execução crítica")))'
                        else:
                            value = etapa_do_status(status)
                            c = f"{get_column_letter(indice_status)}{numero}"
                            formula = f'IF(ISNUMBER(SEARCH("encerrad",{c})),"Encerrados",IF(ISNUMBER(SEARCH("andamento",{c})),"Em andamento","Futuros projetos"))'
                            if values[1] and not str(values[1]).upper().startswith("TOTAL"):
                                origem_grupos.append(value)
                        formulas_novas.setdefault(label, formula)
                        cells[indice] = _texto_cell(ref, value, formula=formula)
            abertura = re.sub(r'\bspans="[^"]+"', f'spans="1:{max(nova_ultima,resumo_fim+deslocamento)}"', match[1])
            return abertura + "".join(v for _, v in sorted(cells.items())) + match[4]

        xml = ROW.sub(revisar_linha, xml)
        xml = re.sub(r'<dimension ref="[^"]+"\s*/>', f'<dimension ref="A1:{get_column_letter(max(nova_ultima,resumo_fim+deslocamento))}{fim}"/>', xml, count=1)
        def ajustar_col(match):
            tag = match[0]
            minimo = int(re.search(r'\bmin="([0-9]+)"', tag)[1])
            maximo = int(re.search(r'\bmax="([0-9]+)"', tag)[1])
            if deslocamento and resumo_inicio <= minimo <= maximo <= resumo_fim:
                tag = re.sub(r'\bmin="[0-9]+"', f'min="{minimo+deslocamento}"', tag)
                tag = re.sub(r'\bmax="[0-9]+"', f'max="{maximo+deslocamento}"', tag)
            elif ultima < minimo <= maximo <= nova_ultima:
                return ""
            return tag
        xml = re.sub(r'<col\b[^>]*/>', ajustar_col, xml)
        cols = "".join(f'<col min="{i}" max="{i}" width="24" customWidth="1"/>' for i in range(ultima+1, nova_ultima+1))
        xml = xml.replace("</cols>", cols+"</cols>")
        xml = re.sub(r'<cols>(.*?)</cols>', lambda m: '<cols>'+''.join(sorted(re.findall(r'<col\b[^>]*/>',m[1]),key=lambda c:int(re.search(r'\bmin="([0-9]+)"',c)[1])))+'</cols>', xml, flags=re.S)
        dv = f'<dataValidation type="list" allowBlank="1" showErrorMessage="1" errorTitle="Status do projeto" error="Selecione uma opção da lista." sqref="{get_column_letter(indice_status)}{inicio+1}:{get_column_letter(indice_status)}{fim-1}"><formula1>StatusProjetos</formula1></dataValidation>'
        if "<dataValidations" in xml:
            raise ValueError("A planilha já contém validações não previstas; nada foi gravado.")
        insertion = re.search(r'<(?:hyperlinks|printOptions|pageMargins|pageSetup|headerFooter|rowBreaks|drawing|tableParts|extLst)\b', xml)
        posicao = insertion.start() if insertion else xml.index("</worksheet>")
        xml = xml[:posicao] + '<dataValidations count="1">'+dv+'</dataValidations>' + xml[posicao:]
        partes[worksheet_path] = xml.encode("utf-8")

        table_xml = partes[tabelas[nome]].decode("utf-8")
        ref = f"{get_column_letter(a)}{inicio}:{get_column_letter(nova_ultima)}{fim}"
        table_xml = table_xml.replace(f'ref="{tabela.ref}"', f'ref="{ref}"')
        for velho, novo in ((headers[final_col-a], "FINAL PREVISTO NA SÍNTESE"), (headers[inicio_col-a] if inicio_col else None, "ASSINATURA DA SÍNTESE")):
            if velho:
                table_xml = table_xml.replace("name="+quoteattr(velho), "name="+quoteattr(novo))
        next_id = max(c.id for c in tabela.tableColumns) + 1
        colunas = "".join(
            f'<tableColumn id="{next_id+i}" name={quoteattr(label)}>'
            +(f'<calculatedColumnFormula>{escape(formulas_novas[label])}</calculatedColumnFormula>' if label in formulas_novas else "")
            +'</tableColumn>' for i,label in enumerate(novos)
        )
        table_xml = re.sub(r'(<tableColumns\b[^>]*\bcount=")[0-9]+', lambda m: m[1]+str(len(headers)+len(novos)), table_xml)
        partes[tabelas[nome]] = table_xml.replace("</tableColumns>", colunas+"</tableColumns>").encode("utf-8")
        if deslocamento:
            recon_xml = partes[tabelas[nome_resumo]].decode("utf-8").replace(f'ref="{resumo.ref}"', f'ref="{_ref_intervalo(resumo.ref,deslocamento)}"')
            recon_xml = re.sub(r'(<calculatedColumnFormula\b[^>]*>)(.*?)(</calculatedColumnFormula>)', lambda m:m[1]+_mover_refs(m[2],resumo_inicio,resumo_fim,deslocamento)+m[3],recon_xml,flags=re.S)
            partes[tabelas[nome_resumo]] = recon_xml.encode("utf-8")
        report.append({"tabela":nome,"intervalo":ref,"reconciliacao":_ref_intervalo(resumo.ref,deslocamento),"etapas":origem_grupos})
    wb.close()

    workbook = partes["xl/workbook.xml"].decode("utf-8")
    rels = partes["xl/_rels/workbook.xml.rels"].decode("utf-8")
    content_types = partes["[Content_Types].xml"].decode("utf-8")
    rid = "rId"+str(max(int(r.attrib["Id"][3:]) for r in relationships if r.attrib["Id"].startswith("rId"))+1)
    sheet_id = max(int(s.attrib["sheetId"]) for s in book.find(f"{{{NS}}}sheets"))+1
    sheet_path = f"xl/worksheets/sheet{sheet_id}.xml"
    workbook = workbook.replace("</sheets>", f'<sheet name="_Opcoes_Projetos" sheetId="{sheet_id}" state="hidden" r:id="{rid}"/></sheets>')
    defined = f'<definedNames><definedName name="StatusProjetos">_Opcoes_Projetos!$A$2:$A${len(statuses)+1}</definedName></definedNames>'
    if "<definedNames" in workbook:
        raise ValueError("Há nomes definidos na planilha; a revisão requer conferência manual.")
    posicao = workbook.find("<calcPr")
    workbook = workbook[:posicao]+defined+workbook[posicao:] if posicao >= 0 else workbook.replace("</workbook>",defined+"</workbook>")
    workbook = re.sub(r'<calcPr\b[^>]*/>', '<calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>', workbook)
    rels = re.sub(r'<Relationship\b[^>]*Type="[^"]*/calcChain"[^>]*/>', "", rels)
    rels = rels.replace("</Relationships>",f'<Relationship Id="{rid}" Type="{REL}/worksheet" Target="worksheets/sheet{sheet_id}.xml"/></Relationships>')
    content_types = re.sub(r'<Override\b[^>]*PartName="/xl/calcChain.xml"[^>]*/>', "", content_types)
    content_types = content_types.replace("</Types>",f'<Override PartName="/{sheet_path}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
    partes["xl/workbook.xml"] = workbook.encode("utf-8")
    partes["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    partes["[Content_Types].xml"] = content_types.encode("utf-8")
    partes.pop("xl/calcChain.xml", None)
    rows = [_texto_cell("A1","Status dos projetos")]
    sheet_rows = ['<row r="1">'+rows[0]+'</row>']
    sheet_rows.extend(f'<row r="{i}">'+_texto_cell(f"A{i}",s)+"</row>" for i,s in enumerate(statuses,start=2))
    partes[sheet_path] = (f'<worksheet xmlns="{NS}"><sheetData>'+"".join(sheet_rows)+"</sheetData></worksheet>").encode("utf-8")
    temporary = NamedTemporaryFile(suffix=".xlsx", dir=caminho.parent, delete=False)
    temporary.close()
    temporario = Path(temporary.name)
    try:
        with ZipFile(temporario,"w",compression=ZIP_DEFLATED) as arquivo:
            for info in infos:
                if info.filename in partes:
                    arquivo.writestr(info,partes.pop(info.filename))
            for nome,dados in partes.items():
                arquivo.writestr(nome,dados)
        depois = parse_financeiro_geral(BytesIO(temporario.read_bytes()))
        if depois["erros"]:
            raise ValueError("Revisão inválida: " + " | ".join(depois["erros"]))
        for chave in ("resumos","projetos"):
            novos_campos = {"inicio"} if chave == "projetos" else set()
            novos_campos |= {"data_kickoff","fim_previsto_real","etapa","farol"} if chave == "projetos" else set()
            limpar = lambda itens: [{k:v for k,v in item.items() if k not in novos_campos} for item in itens]
            if limpar(antes[chave]) != limpar(depois[chave]):
                raise ValueError(f"A revisão alteraria dados financeiros de {chave}; nada foi gravado.")
        # Datas de assinatura existentes também precisam permanecer idênticas.
        if [r["inicio"] for r in antes["projetos"]] != [r["inicio"] for r in depois["projetos"]]:
            raise ValueError("As datas existentes foram alteradas; nada foi gravado.")
        try:
            temporario.replace(caminho)
        except PermissionError:
            # Pré-visualizadores do Windows podem permitir edição e bloquear
            # a troca do arquivo. Os bytes já foram validados e há backup.
            revisado = temporario.read_bytes()
            gravacao_iniciada = False
            try:
                with caminho.open("r+b") as destino:
                    gravacao_iniciada = True
                    destino.write(revisado)
                    destino.truncate()
                    destino.flush()
                    os.fsync(destino.fileno())
                if caminho.read_bytes() != revisado:
                    raise OSError("Não foi possível conferir os bytes gravados.")
            except Exception:
                if gravacao_iniciada:
                    caminho.write_bytes(original)
                if backup_dir is not None:
                    (backup_dir / (caminho.stem + " - revisado.xlsx")).write_bytes(revisado)
                raise
    finally:
        if temporario.exists():
            temporario.unlink()
    result = {"alterado":True,"sha256_antes":assinatura,"sha256_depois":sha256(caminho.read_bytes()).hexdigest(),"projetos":len(depois["projetos"]),"tabelas":report}
    if backup_dir is not None:
        (backup_dir/"verificacao-refat-v2.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return result
