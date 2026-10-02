"""Leitura do FINANCEIRO GERAL REFAT por objetos Table e cabeçalhos.

As origens TAB.* preservam o contrato dos snapshots existentes. Nenhuma
posição de célula, nome de aba ou tamanho de tabela identifica os dados.
"""
from __future__ import annotations

import re
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel

from .comum import (
    abrir_workbook, normalizar, novo_payload,
    para_decimal, para_dinheiro, texto_limpo,
)


@dataclass(frozen=True)
class Campo:
    cabecalhos: tuple[str, ...]
    tipo: str = "dinheiro"
    opcional: bool = False


@dataclass(frozen=True)
class TabelaFinanceira:
    origem: str
    campos: dict[str, Campo]
    projetos: bool = False
    pilar: str | None = None


def _campo(*cabecalhos, tipo="dinheiro", opcional=False):
    return Campo(cabecalhos, tipo, opcional)


_COMBINADO = _campo("REALIZADO + PROJETADO")
_PROJETADO_TOTAL = _campo("PROJETADO TOTAL")
_PROJETOS = {
    "nome": _campo("PROJETO", tipo="texto"),
    "status": _campo("STATUS", tipo="texto"),
    "sequencia": _campo("ITEM", tipo="inteiro", opcional=True),
    "inicio": _campo("INICIO", "INICIO DO PROJETO", tipo="data"),
    "fim": _campo("FINAL", "FINAL DO PROJETO", "DATA FINAL DO PROJETO", tipo="data"),
    "orcado": _campo("ORÇADO SÍNTESE", "ORÇADO ORIGINAL", "VALOR ORÇADO"),
    "realizado": _campo("REALIZADO", "REALIZADO TOTAL"),
    "projetado_2026": _campo("PROJETADO 2026"),
    "projetado_2027": _campo("PROJETADO 2027"),
    "realizado_mais_projetado": _COMBINADO,
    "diferenca": _campo("DIFERENÇA ORÇADO"),
    "percentual": _campo("% REALIZADO ORÇADO", "% REALIZADO", tipo="percentual"),
}
_RECONCILIACAO = {
    "realizado": _campo("REALIZADO"),
    "projetado": _PROJETADO_TOTAL,
    "realizado_mais_projetado": _COMBINADO,
    "diferenca": _campo("DIFERENÇA PPI", "DIF. REL. AO PPI"),
}
TABELAS_FINANCEIRO = {
    "tbl_ConsolidadoPrograma": TabelaFinanceira("TAB. 1", {
        "rotulo": _campo("AÇÃO", tipo="texto"),
        "recurso_ou_meta": _campo("RECURSO PPI TOTAL"),
        "realizado": _campo("REALIZADO TOTAL"),
        "projetado": _PROJETADO_TOTAL,
        "realizado_mais_projetado": _COMBINADO,
        "diferenca": _campo("DIFERENÇA PPI"),
        "percentual": _campo("% EXECUÇÃO", tipo="percentual"),
        "farol": _campo("FAROL", tipo="texto"),
    }),
    "tbl_VisaoGeral": TabelaFinanceira("TAB. 1.1", {
        "rotulo": _campo("AÇÃO", tipo="texto"),
        "recurso_ou_meta": _campo("RECURSO PPI TOTAL"),
        "realizado_2024": _campo("REALIZADO 2024"),
        "realizado_2025": _campo("REALIZADO 2025"),
        "realizado_2026": _campo("REALIZADO 2026"),
        "projetado_2026": _campo("PROJETADO 2026"),
        "projetado_2027": _campo("PROJETADO 2027"),
        "projetado": _PROJETADO_TOTAL,
        "realizado_mais_projetado": _COMBINADO,
        "diferenca": _campo("DIFERENÇA PPI"),
        "percentual": _campo("% EXECUÇÃO", tipo="percentual"),
        "farol": _campo("FAROL", tipo="texto"),
    }),
    "tbl_ProjetosPDI": TabelaFinanceira("TAB. 2", _PROJETOS, True, "PDI"),
    "tbl_VisaoPDI": TabelaFinanceira("TAB. 3", {
        **_RECONCILIACAO,
        "recurso_ou_meta": _campo("RECURSO PPI AFCCT / PDI"),
    }, pilar="PDI"),
    "tbl_ProjetosFCRH": TabelaFinanceira("TAB. 4", _PROJETOS, True, "FORMACAO"),
    "tbl_VisaoFCRH": TabelaFinanceira("TAB. 5", {
        **_RECONCILIACAO,
        "recurso_ou_meta": _campo("RECURSO PPI FCRH"),
    }, pilar="FORMACAO"),
    "tbl_ProjetosACS": TabelaFinanceira("TAB. 6", {
        **_PROJETOS, "inicio": _campo("INICIO", tipo="data", opcional=True),
    }, True, "STARTUPS"),
    "tbl_VisaoACS": TabelaFinanceira("TAB. 8", {
        **_RECONCILIACAO, "recurso_ou_meta": _campo("TOTAL ACS"),
    }, pilar="STARTUPS"),
    "tbl_CaptaoATeOutros": TabelaFinanceira("TAB. 9", {
        "rotulo": _campo("PILAR", tipo="texto"),
        "recurso_ou_meta": _campo("META TOTAL"),
        "captado_anos_1_2": _campo("CAPTADO (ANO 1 E 2)"),
        "captado_ano3_ytd": _campo("CAPTADO (ANO 3 - YTD)"),
        "captado": _campo("CAPTADO TOTAL"),
        "realizado": _campo("REALIZADO TOTAL"),
        "percentual_meta": _campo("% CAPTADO META", tipo="percentual"),
        "percentual_captado": _campo("% REALIZADO CAPTADO", tipo="percentual"),
    }),
}
TABELAS_OBRIGATORIAS = tuple(TABELAS_FINANCEIRO)
_PILARES = {
    "afcct/pd&i": "PDI", "afcct/pdi": "PDI", "fcrh": "FORMACAO",
    "acs": "STARTUPS", "infraestrutura": "INFRA",
    "associacaotecnologica": "AT", "outrasfontes": "OUTRASFONTES",
}
_TOTAIS = {"total", "total geral", "total geral - atual"}
_CAMPOS_RESUMO = (
    "recurso_ou_meta", "captado_anos_1_2", "captado_ano3_ytd", "captado",
    "realizado", "projetado", "realizado_mais_projetado", "diferenca",
    "percentual", "percentual_meta", "percentual_captado", "farol",
)
_CAMPOS_PROJETO = (
    "sequencia", "nome", "status", "inicio", "fim", "orcado", "realizado",
    "projetado_2026", "projetado_2027", "realizado_mais_projetado",
    "diferenca", "percentual",
)
_LIMITES_TEXTO = {"nome": 255, "status": 120, "farol": 20}


def _cabecalho(valor):
    return normalizar(re.sub(r"_x000a_", "\n", texto_limpo(valor), flags=re.I))


def _chave_campo(valor):
    texto = _cabecalho(valor)
    # PDI identifica o ano apenas no intervalo de meses do cabeçalho.
    periodo = re.fullmatch(r"projetado\s*\([^)]*/(\d{4})\)", texto)
    if periodo:
        return f"projetado {periodo[1]}"
    # Apenas qualificadores temporais conhecidos; captação conserva ANO 1/2/3.
    texto = re.sub(r"\s*\((?:ate\s+[^)]*|\d{2}\s+a\s+\d{2}/\d{4}|ytd)\)", "", texto)
    return re.sub(r"\s+ate\s+[\w/]+$", "", texto).strip()


def catalogar_tabelas(wb):
    """Nome do objeto → (worksheet, Table); inclui todas as abas."""
    catalogo = {}
    for ws in wb.worksheets:
        for tabela in ws.tables.values():
            if tabela.displayName in catalogo:
                raise ValueError(f'Tabela duplicada: "{tabela.displayName}".')
            catalogo[tabela.displayName] = (ws, tabela)
    return catalogo


def parse_financeiro_geral(caminho: Path | str) -> dict:
    """Parser puro do REFAT, mantendo o payload e as origens dos snapshots."""
    payload = novo_payload()
    with ExitStack() as recursos:
        wb = abrir_workbook(caminho, read_only=False)
        recursos.callback(wb.close)
        try:
            catalogo = catalogar_tabelas(wb)
        except ValueError as exc:
            payload["erros"].append(str(exc))
            return payload
        for nome in TABELAS_OBRIGATORIAS:
            if nome not in catalogo:
                payload["erros"].append(
                    f'Tabela obrigatória "{nome}" não encontrada. Envie o modelo FINANCEIRO GERAL - REFAT.xlsx.')
        formulas = abrir_workbook(caminho, data_only=False)
        recursos.callback(formulas.close)
        chaves = set()
        for nome, especificacao in TABELAS_FINANCEIRO.items():
            if nome in catalogo:
                ws, tabela = catalogo[nome]
                _ler_tabela(payload, ws, tabela, formulas[ws.title], especificacao, wb.epoch, chaves)
    if "tbl_VisaoGeral" in catalogo and not any(r["ano"] is not None for r in payload["resumos"]):
        payload["avisos"].append("tbl_VisaoGeral sem valores anuais por pilar; a visão anual exibirá —.")
    return payload


def _mapear_colunas(payload, ws, tabela, especificacao, limites):
    primeira, inicio, ultima, _fim = limites
    contexto = f'Aba "{ws.title}", tabela "{tabela.displayName}"'
    if tabela.headerRowCount != 1 or len(tabela.tableColumns) != ultima - primeira + 1:
        payload["erros"].append(f"{contexto}: cabeçalhos ou quantidade de colunas inválidos.")
        return None
    nomes = [coluna.name for coluna in tabela.tableColumns]
    cabecalhos = [ws.cell(inicio, coluna).value for coluna in range(primeira, ultima + 1)]
    if [_cabecalho(v) for v in nomes] != [_cabecalho(v) for v in cabecalhos]:
        payload["erros"].append(f"{contexto}: cabeçalho divergente dos metadados da tabela.")
        return None
    normalizados = [_cabecalho(v) for v in nomes]
    if len(set(normalizados)) != len(normalizados):
        payload["erros"].append(f"{contexto}: cabeçalhos duplicados após normalização.")
        return None
    mapa = {}
    problemas = []
    for campo, contrato in especificacao.campos.items():
        aliases = {_chave_campo(v) for v in contrato.cabecalhos}
        indices = [i for i, nome in enumerate(nomes) if _chave_campo(nome) in aliases]
        if len(indices) > 1:
            problemas.append(f'cabeçalho ambíguo para "{campo}"')
        elif indices:
            mapa[campo] = indices[0]
        elif not contrato.opcional:
            problemas.append(f'coluna obrigatória "{contrato.cabecalhos[0]}" ausente')
    if problemas:
        payload["erros"].append(f"{contexto}: {'; '.join(problemas)}.")
        return None
    return mapa


def _ler_tabela(payload, ws, tabela, formulas, especificacao, epoch, chaves):
    try:
        limites = range_boundaries(tabela.ref)
    except (ValueError, TypeError):
        payload["erros"].append(f'Tabela "{tabela.displayName}": intervalo inválido.')
        return
    mapa = _mapear_colunas(payload, ws, tabela, especificacao, limites)
    if mapa is None:
        return
    primeira, inicio, ultima, fim = limites
    marcadores = [i for i, coluna in enumerate(tabela.tableColumns)
                  if _cabecalho(coluna.name) in {"acao", "pilar", "projeto", "item", "coluna1"}]
    intervalo = dict(min_col=primeira, max_col=ultima, min_row=inicio + 1,
                     max_row=fim - (tabela.totalsRowCount or 0))
    if intervalo["min_row"] > intervalo["max_row"]:
        return
    for linha, linha_formula in zip(ws.iter_rows(**intervalo), formulas.iter_rows(**intervalo), strict=True):
        if any(_cabecalho(linha[i].value) in _TOTAIS for i in marcadores):
            continue
        if all(linha[i].value is None and linha_formula[i].data_type != "f" for i in mapa.values()):
            continue
        valores = {
            campo: _converter(payload, celula=linha[indice], formula=linha_formula[indice],
                              campo=campo, tipo=especificacao.campos[campo].tipo,
                              aba=ws.title, tabela=tabela.displayName, epoch=epoch)
            for campo, indice in mapa.items()
        }
        pilar = especificacao.pilar or _PILARES.get(normalizar(valores.get("rotulo")).replace(" ", ""))
        identidade = valores.get("nome") if especificacao.projetos else pilar
        coluna_identidade = "nome" if especificacao.projetos else "rotulo"
        coordenada = linha[mapa.get(coluna_identidade, 0)].coordinate
        contexto = f'Aba "{ws.title}", tabela "{tabela.displayName}", célula {coordenada}'
        if not identidade:
            problema = "nome do projeto vazio" if especificacao.projetos else f'pilar desconhecido "{valores.get("rotulo") or ""}"'
            payload["erros"].append(f"{contexto}: {problema}.")
            continue
        chave = (especificacao.origem, identidade)
        if chave in chaves:
            payload["erros"].append(f'{contexto}: registro duplicado "{identidade}".')
            continue
        chaves.add(chave)
        campos = _CAMPOS_PROJETO if especificacao.projetos else _CAMPOS_RESUMO
        registro = {campo: valores.get(campo) for campo in campos}
        registro.update(origem=especificacao.origem, pilar=pilar, ano=None)
        payload["projetos" if especificacao.projetos else "resumos"].append(registro)
        if especificacao.origem == "TAB. 1.1":
            for ano in (2024, 2025, 2026, 2027):
                realizado = valores.get(f"realizado_{ano}")
                projetado = valores.get(f"projetado_{ano}")
                if realizado is not None or projetado is not None:
                    anual = {campo: None for campo in _CAMPOS_RESUMO}
                    anual.update(origem=especificacao.origem, pilar=pilar, ano=ano,
                                 realizado=realizado, projetado=projetado)
                    payload["resumos"].append(anual)


def _converter(payload, *, celula, formula, campo, tipo, aba, tabela, epoch):
    bruto = celula.value
    contexto = f'Aba "{aba}", tabela "{tabela}", campo "{campo}", célula {celula.coordinate}'
    if celula.data_type == "e":
        payload["avisos"].append(f"{contexto}: erro de fórmula {bruto}; valor indisponível.")
        return None
    if bruto is None:
        if formula.data_type == "f":
            payload["avisos"].append(f"{contexto}: fórmula sem resultado salvo; recalcule e salve no Excel.")
        return None
    if tipo == "texto":
        texto = texto_limpo(bruto)
        if len(texto) > _LIMITES_TEXTO.get(campo, 255):
            payload["erros"].append(f"{contexto}: texto excede o tamanho permitido.")
        return texto
    if texto_limpo(bruto) in {"", "-", "–", "—"}:
        return None
    try:
        if tipo == "data":
            if isinstance(bruto, datetime):
                return bruto.date()
            if isinstance(bruto, date):
                return bruto
            if isinstance(bruto, (int, float)) and not isinstance(bruto, bool):
                convertido = from_excel(bruto, epoch)
                if isinstance(convertido, datetime):
                    return convertido.date()
            if isinstance(bruto, str):
                for formato in ("%Y-%m-%d", "%d/%m/%Y"):
                    try:
                        return datetime.strptime(bruto.strip(), formato).date()
                    except ValueError:
                        pass
            raise ValueError("data inválida")
        numero = para_dinheiro(bruto) if tipo == "dinheiro" else para_decimal(bruto)
        if numero is None or not numero.is_finite():
            raise ValueError("valor não numérico")
        if tipo == "inteiro":
            if numero != numero.to_integral_value() or not 0 <= numero <= 2147483647:
                raise ValueError("sequência deve ser um inteiro não negativo")
            return int(numero)
        limite = Decimal("1e16") if tipo == "dinheiro" else Decimal("1e4")
        if abs(numero) >= limite:
            raise ValueError("valor excede o tamanho permitido")
        return numero
    except (ValueError, TypeError, InvalidOperation, OverflowError) as exc:
        payload["erros"].append(f'{contexto}: {exc} (valor "{texto_limpo(bruto)}").')
        return None
