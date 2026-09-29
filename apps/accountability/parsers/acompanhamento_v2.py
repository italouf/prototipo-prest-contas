"""Parser de ``Acompanhamento Financeiro (v2).xlsx`` — abas 3–9 (SDD §6).

Metadados por *rótulo* (não por coordenada fixa): ``0. Sumário`` é a fonte
primária e os valores lidos nas abas 3–9 são conferidos contra ela —
divergência vira erro citando aba e célula. As colunas das abas de despesa
são mapeadas pelas 23 regras ordenadas do SDD §6 (primeira casa vence) sobre
o texto normalizado do cabeçalho; cabeçalhos não mapeados são ignorados.

REGRA DURA: ``6.1 Conta Ação - AT (Lei TICs)`` é segregada como
``AT_LEI_TICS`` e nunca soma em ``AT``; ``9. Ampliação de Infraestrutura`` é
aquisição patrimonial (sem data de pagamento esperada). A aba ``10. Equipe``
e as demais não financeiras nunca são lidas.

O parser é puro: lê o ``.xlsx`` com ``openpyxl`` (``read_only``,
``data_only``) e devolve o payload de ``comum.novo_payload()``; nunca toca
no banco. Por desempenho, a varredura de cada aba cobre no máximo
``_MAX_COLUNAS`` colunas e ``_MAX_LINHAS_DADOS`` linhas de dados.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from apps.accountability.parsers.comum import (
    abrir_workbook,
    eh_erro_formula,
    normalizar,
    novo_payload,
    para_decimal,
    para_dinheiro,
    texto_limpo,
)

# Aba → (pilar, tipo_recurso) — SDD §6. 6.1 nunca vira AT.
ABAS_DESPESA: dict[str, tuple[str, str]] = {
    "3. Conta Ação - AFCCT": ("PDI", "EMBRAPII"),
    "4. Conta Ação - FCRH": ("FORMACAO", "EMBRAPII"),
    "5. Conta Ação - ACS": ("STARTUPS", "EMBRAPII"),
    "6. Conta Ação - AT": ("AT", "AT"),
    "6.1 Conta Ação - AT (Lei TICs)": ("AT", "AT_LEI_TICS"),
    "7. Conta Ação - Outras Fontes": ("OUTRASFONTES", "OUTRAS_FONTES"),
    "8. Conta Ação - Infraestrutura": ("INFRA", "EMBRAPII"),
    "9. Ampliação de Infraestrutura": ("INFRA", "EMBRAPII"),
}

ABA_SUMARIO = "0. Sumário"

# Limites de varredura (desempenho): sheet "5. Conta Ação - ACS" declara
# max_column=16383 por formatação — varrer a linha inteira é patológico.
_MAX_COLUNAS = 60
_MAX_LINHAS_DADOS = 5000
_MAX_LINHA_CABECALHO = 20
_MAX_LINHA_METADADOS = 12
_MAX_LINHA_METADADOS_SUMARIO = 40  # no Sumário real os rótulos ficam em B19:B21

# Rótulo normalizado → chave do payload de metadados.
_ROTULOS_METADADOS = (
    ("centro de competencia", "centro"),
    ("termo de cooperacao", "termo"),
    ("periodo de referencia", "periodo_referencia"),
)

_MARCADORES_VAZIOS = ("-", "–", "—")

_CAMPOS_DATA = ("data_nota", "data_pagamento", "data_movimento")
_CAMPOS_NUMEROS = ("quantidade", "valor_unitario", "valor")
_CAMPOS_DESPESA = (
    "sequencia", "acao_relacionada", "codigo_projeto", "marco", "tipo_despesa",
    "credor", "documento", "numero_nota", "link_documento", "numero_patrimonial",
    "descricao", "descricao_atividade", "observacao", "fonte_recurso",
    "data_nota", "data_pagamento", "data_movimento",
    "quantidade", "valor_unitario", "valor",
)


def parse_acompanhamento_v2(caminho: Path | str) -> dict:
    """Lê ``Acompanhamento Financeiro (v2).xlsx`` e devolve o payload padrão.

    O payload tem o mesmo shape de ``comum.PAYLOAD_VAZIO``. Falha de leitura
    do arquivo propaga ``comum.ArquivoInvalidoError``.
    """
    payload = novo_payload()
    wb = abrir_workbook(caminho)
    try:
        presentes = [aba for aba in ABAS_DESPESA if aba in wb.sheetnames]
        if not presentes:
            payload["erros"].append(
                "Não foi possível identificar o tipo do arquivo: nenhuma das abas"
                " do Acompanhamento Financeiro (v2) foi encontrada"
                f" (esperado: {', '.join(repr(a) for a in ABAS_DESPESA)})."
            )
            return payload
        for aba in ABAS_DESPESA:
            if aba not in wb.sheetnames:
                payload["erros"].append(
                    f'A aba "{aba}" não foi encontrada.'
                    " Verifique se é o template correto da EMBRAPII."
                )
        _ler_metadados(wb, payload)
        for aba in presentes:
            _ler_aba_despesa(payload, wb[aba], aba)
    finally:
        wb.close()
    return payload


# ----------------------------------------------------------------- metadados

def _ler_metadados(wb, payload: dict) -> None:
    """Metadados por rótulo: ``0. Sumário`` primeiro, abas 3–9 conferidas.

    Divergência entre a fonte primária e uma aba vira erro citando aba e a
    célula onde o valor divergente foi lido (SDD §6/§7).
    """
    origem: dict[str, tuple[str, str, str]] = {}  # chave -> (valor, aba, coordenada)
    if ABA_SUMARIO in wb.sheetnames:
        for chave, valor, coordenada in _varrer_metadados(
                wb[ABA_SUMARIO], _MAX_LINHA_METADADOS_SUMARIO):
            origem.setdefault(chave, (valor, ABA_SUMARIO, coordenada))
    for aba in ABAS_DESPESA:
        if aba not in wb.sheetnames:
            continue
        for chave, valor, coordenada in _varrer_metadados(wb[aba], _MAX_LINHA_METADADOS):
            if chave not in origem:
                origem[chave] = (valor, aba, coordenada)
                continue
            referencia, aba_fonte, _ = origem[chave]
            if normalizar(referencia) != normalizar(valor):
                payload["erros"].append(
                    f"Metadado divergente em '{aba}' ({coordenada}):"
                    f" {_rotulo_de(chave)} \"{valor}\" difere de \"{referencia}\""
                    f" ({aba_fonte})."
                )
    for chave, (valor, _aba, _coordenada) in origem.items():
        payload["metadados"][chave] = valor


def _varrer_metadados(ws, limite: int) -> list[tuple[str, str, str]]:
    """``(chave, valor, coordenada)`` por rótulo nas primeiras ``limite`` linhas.

    O valor é a próxima célula não vazia à direita do rótulo e a varredura
    segue depois dela: o valor de "Centro de Competência ..." também casa o
    prefixo do rótulo e não pode ser lido como rótulo.
    """
    achados: list[tuple[str, str, str]] = []
    linhas = ws.iter_rows(min_row=1, max_row=limite, min_col=1, max_col=_MAX_COLUNAS)
    for numero, linha in enumerate(linhas, start=1):
        indice = 0
        while indice < len(linha):
            chave = _chave_metadado(_valor_da_celula(linha[indice]))
            if chave is None:
                indice += 1
                continue
            valor_indice = None
            for j in range(indice + 1, len(linha)):
                if _tem_valor(_valor_da_celula(linha[j])):
                    valor_indice = j
                    break
            if valor_indice is None:
                indice += 1
                continue
            achados.append((
                chave,
                texto_limpo(_valor_da_celula(linha[valor_indice])),
                f"{get_column_letter(valor_indice + 1)}{numero}",
            ))
            indice = valor_indice + 1
    return achados


def _chave_metadado(valor: object) -> str | None:
    texto = normalizar(valor)
    for prefixo, chave in _ROTULOS_METADADOS:
        if texto.startswith(prefixo):
            return chave
    return None


def _rotulo_de(chave: str) -> str:
    for prefixo, atual in _ROTULOS_METADADOS:
        if atual == chave:
            return prefixo
    return chave


# ------------------------------------------------------------- abas de despesa

def _ler_aba_despesa(payload: dict, ws, aba: str) -> None:
    """Lê o cabeçalho (linha com célula ``Linha``) e as linhas de despesa.

    A leitura **encerra** na primeira linha em que nenhum campo mapeado (além
    de ``sequencia``) tenha valor — a linha do template com só o número
    sequencial termina a leitura (SDD §6).
    """
    pilar, tipo_recurso = ABAS_DESPESA[aba]
    erros = payload["erros"]
    cabecalho: int | None = None
    colunas: dict[str, int] = {}  # campo -> índice 0-based (primeira coluna vence)
    linhas = ws.iter_rows(
        min_row=1, max_row=_MAX_LINHA_CABECALHO + _MAX_LINHAS_DADOS + 1,
        min_col=1, max_col=_MAX_COLUNAS)
    lidas = 0
    for numero, linha in enumerate(linhas, start=1):
        if cabecalho is None:
            if numero <= _MAX_LINHA_CABECALHO:
                indice = _indice_coluna_sequencia(linha)
                if indice is not None:
                    cabecalho = numero
                    colunas = _mapear_colunas(linha)
                continue
            erros.append(
                f'A aba "{aba}": cabeçalho "Linha" não encontrado nas primeiras'
                f" {_MAX_LINHA_CABECALHO} linhas (linha de cabeçalho esperada até"
                " a linha 20). Verifique se é o template correto da EMBRAPII."
            )
            return
        if numero <= cabecalho:
            continue
        if not any(_tem_valor(_valor_da_celula(linha[idx]))
                   for campo, idx in colunas.items() if campo != "sequencia"):
            # Linha do template com só a sequência (ou vazia): encerra a leitura.
            # Erro de fórmula em coluna não mapeada vira aviso, nunca erro (SDD §7).
            _avisar_formulas_nao_mapeadas(payload, aba, numero, linha, colunas)
            return
        lidas += 1
        if lidas > _MAX_LINHAS_DADOS:
            payload["avisos"].append(
                f'A aba "{aba}" tem mais de {_MAX_LINHAS_DADOS} linhas de dados;'
                " leitura interrompida."
            )
            return
        despesa = _montar_despesa(payload, aba, numero, linha, colunas,
                                  pilar, tipo_recurso)
        payload["despesas"].append(despesa)


def _montar_despesa(
    payload: dict,
    aba: str,
    numero: int,
    linha: tuple,
    colunas: dict[str, int],
    pilar: str,
    tipo_recurso: str,
) -> dict:
    """Monta o dict da despesa validando erros de fórmula e o campo ``valor``.

    Erro de fórmula em coluna **mapeada** vira erro citando aba, linha e
    coluna (SDD §7); em coluna não mapeada vira aviso e nunca bloqueia — o
    template real traz ``#REF!`` na coluna ``Ano`` (não mapeada) da aba 8
    (Ruling 14).
    """
    erros = payload["erros"]
    mapeadas = set(colunas.values())
    for indice, celula in enumerate(linha):
        bruto = _valor_da_celula(celula)
        if indice in mapeadas and eh_erro_formula(bruto):
            erros.append(
                f'{aba}: linha {numero}, coluna {get_column_letter(indice + 1)}:'
                f' erro de fórmula "{texto_limpo(bruto)}".'
            )
    _avisar_formulas_nao_mapeadas(payload, aba, numero, linha, colunas)

    despesa = {"aba": aba, "linha": numero, "pilar": pilar,
               "tipo_recurso": tipo_recurso}
    despesa.update({campo: None for campo in _CAMPOS_DESPESA})
    for campo, indice in colunas.items():
        bruto = _valor_da_celula(linha[indice])
        if not _tem_valor(bruto) or eh_erro_formula(bruto):
            continue
        coluna = get_column_letter(indice + 1)
        if campo == "valor":
            convertido = para_dinheiro(bruto)
            if convertido is None:
                erros.append(
                    f'{aba}: linha {numero}, coluna {coluna}:'
                    f' valor inválido "{texto_limpo(bruto)}".'
                )
            despesa[campo] = convertido
        elif campo in ("quantidade", "valor_unitario"):
            despesa[campo] = para_decimal(bruto)  # não é moeda: sem quantização
        elif campo == "sequencia":
            numero_seq = para_decimal(bruto)
            despesa[campo] = int(numero_seq) if numero_seq is not None else None
        elif campo in _CAMPOS_DATA:
            despesa[campo] = _para_data(bruto)
        else:
            despesa[campo] = texto_limpo(bruto)
    return despesa


def _mapear_colunas(linha: tuple) -> dict[str, int]:
    """Campo → índice 0-based pelas 23 regras do SDD §6 (primeira casa vence)."""
    colunas: dict[str, int] = {}
    for indice, celula in enumerate(linha):
        texto = normalizar(_valor_da_celula(celula))
        if not texto:
            continue
        campo = _campo_do_cabecalho(texto)
        if campo is not None:
            colunas.setdefault(campo, indice)  # primeira coluna vence
    return colunas


def _campo_do_cabecalho(texto: str) -> str | None:
    """Aplica as 23 regras de coluna do SDD §6 em ordem (primeira casa vence)."""
    if texto == "linha":
        return "sequencia"                                   # 1
    if "codigo do projeto" in texto or "conta do projeto" in texto:
        return "codigo_projeto"                              # 2
    if "data do pagamento" in texto:
        return "data_pagamento"                              # 3
    if texto.startswith("data da nota"):
        return "data_nota"                                   # 4
    if texto == "data":
        return "data_movimento"                              # 5
    if "valor unit" in texto:
        return "valor_unitario"                              # 6
    if "quantidade" in texto:
        return "quantidade"                                  # 7
    if "numero patrimonial" in texto:
        return "numero_patrimonial"                          # 8
    if "numero da nota" in texto:
        return "numero_nota"                                 # 9
    if texto.startswith("valor"):
        return "valor"                                       # 10
    if "link" in texto:
        return "link_documento"                              # 11
    if "cnpj" in texto or "cpf" in texto:
        return "documento"                                   # 12
    if "fornecedor" in texto and "estrangeiro" not in texto:
        return "credor"                                      # 13
    if "entidade que realizou" in texto:
        return "credor"                                      # 14
    if "credor" in texto and "estrangeiro" not in texto:
        return "credor"                                      # 15
    if "tipo de despesa" in texto:
        return "tipo_despesa"                                # 16
    if "marco" in texto:
        return "marco"                                       # 17
    if "acao esta relacionad" in texto:  # radical: o template real usa o masculino
        return "acao_relacionada"                            # 18
    if "fonte recurso" in texto:
        return "fonte_recurso"                               # 19
    if "breve descritivo" in texto:
        return "descricao_atividade"                         # 20
    if "breve descricao" in texto:
        return "descricao"                                   # 21
    if "descricao" in texto:
        return "descricao"                                   # 22
    if "observa" in texto:
        return "observacao"                                  # 23
    return None


def _indice_coluna_sequencia(linha: tuple) -> int | None:
    """Índice da célula cujo texto normalizado é ``linha`` (cabeçalho)."""
    for indice, celula in enumerate(linha):
        if normalizar(_valor_da_celula(celula)) == "linha":
            return indice
    return None


def _avisar_formulas_nao_mapeadas(
    payload: dict, aba: str, numero: int, linha: tuple, colunas: dict[str, int]
) -> None:
    """Erro de fórmula em coluna não mapeada vira aviso, nunca erro (SDD §7)."""
    mapeadas = set(colunas.values())
    for indice, celula in enumerate(linha):
        bruto = _valor_da_celula(celula)
        if indice in mapeadas or not eh_erro_formula(bruto):
            continue
        payload["avisos"].append(
            f'{aba}: linha {numero}, coluna {get_column_letter(indice + 1)}:'
            f' erro de fórmula "{texto_limpo(bruto)}" em coluna não mapeada;'
            " ignorado."
        )


# ----------------------------------------------------------------- conversões

def _valor_da_celula(celula: Any) -> object:
    """``.value`` da célula (``EmptyCell`` do modo read-only não tem atributo)."""
    return getattr(celula, "value", None)


def _tem_valor(bruto: object) -> bool:
    """``True`` para conteúdo real; vazio e marcadores ``-`` contam como vazio."""
    if bruto is None:
        return False
    if isinstance(bruto, str):
        texto = bruto.strip()
        return texto != "" and texto not in _MARCADORES_VAZIOS
    return True


def _para_data(bruto: object) -> date | None:
    """``date`` | ``datetime`` | serial Excel | ``dd/mm/aaaa`` | ``aaaa-mm-dd``.

    Valor inválido ⇒ ``None``, sem erro (campo opcional).
    """
    if isinstance(bruto, datetime):
        return bruto.date()
    if isinstance(bruto, date):
        return bruto
    if isinstance(bruto, bool):
        return None
    if isinstance(bruto, (int, float)):
        try:
            convertido = from_excel(bruto)
        except (ValueError, TypeError, OverflowError):
            return None
        return convertido.date() if isinstance(convertido, datetime) else convertido
    if isinstance(bruto, str):
        texto = bruto.strip()
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(texto, formato).date()
            except ValueError:
                continue
    return None
