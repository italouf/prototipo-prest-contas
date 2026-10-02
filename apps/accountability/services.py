"""Ingestão atômica e idempotente dos arquivos de prestação de contas (SDD §2).

Fluxo (SDD §2): detectar o tipo pelas tabelas/abas do arquivo → parser
puro → validar erros e metadados → buscar/criar ``CentroCompetencia`` e
``Acompanhamento`` → ``transaction.atomic()`` apagando só os fatos daquele
``tipo_fonte`` e recriando-os → ``ImportacaoAcompanhamento`` (SUCESSO|ERRO) +
auditoria.

Regras duras:
- Idempotência: reenviar o mesmo arquivo substitui os fatos da fonte; nunca duplica.
- Isolamento por fonte: reimportar uma fonte não toca nas demais.
- Falha ⇒ rollback total; o ``ImportacaoAcompanhamento`` de ERRO é gravado
  **fora** da transação que falhou (SDD §7).
- Ruling 7: ``KpiAcompanhamento`` é atualizado por ``(acompanhamento, codigo)``
  — nunca apagado e recriado — para que o ``on_delete=CASCADE`` de
  ``OverrideAcompanhamento.kpi`` não destrua os overrides manuais (SDD §9).
- Ruling 5: a unicidade de ``ResumoFinanceiro`` não vale para ``ano=NULL``; o
  delete do escopo inteiro da fonte antes do insert é o que impede duplicatas.
- Valores vazios viram ``NULL`` (numéricos) ou ``""`` (texto ``blank``); nunca
  ``0`` implícito. Nada do Excel é executado — só lido.
- Ruling 22: v2 sem metadados legíveis importa com os valores do formulário e
  um aviso de que a validação contra o arquivo não foi possível (nunca erro).
"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from openpyxl.worksheet.table import Table
from openpyxl.xml.functions import fromstring

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.text import slugify

from apps.audit.services import registrar_auditoria
from apps.pillars.models import Pilar

from .models import (
    Acompanhamento,
    CentroCompetencia,
    DespesaAcompanhamento,
    ImportacaoAcompanhamento,
    KpiAcompanhamento,
    ProjetoFinanceiro,
    ResumoFinanceiro,
)
from .parsers import acompanhamento_v2, financeiro_geral, indicadores_pe
from .parsers.comum import ArquivoInvalidoError, abrir_workbook, normalizar


class IngestaoError(Exception):
    """Rejeição de upload com mensagem amigável em pt-BR (SDD §7)."""


_PARSERS = {
    "FINANCEIRO_GERAL": financeiro_geral.parse_financeiro_geral,
    "INDICADORES_PE": indicadores_pe.parse_indicadores_pe,
    "ACOMPANHAMENTO_V2": acompanhamento_v2.parse_acompanhamento_v2,
}

# Financeiro: objetos Table. Outras fontes: abas obrigatórias.
VALIDADORES: dict[str, tuple[str, ...]] = {
    "FINANCEIRO_GERAL": financeiro_geral.TABELAS_OBRIGATORIAS,
    "INDICADORES_PE": ("HEAD - INDICADORES - EMBRAPII",),
    "ACOMPANHAMENTO_V2": tuple(acompanhamento_v2.ABAS_DESPESA),
}

# Assinaturas de detecção pelas abas (mesma semântica dos parsers).
_ASSINATURA_INDICADORES = "head - indicadores"
_ASSINATURA_V2 = {normalizar(a) for a in acompanhamento_v2.ABAS_DESPESA} | {
    normalizar(acompanhamento_v2.ABA_SUMARIO)
}

# tipo_fonte → lista do payload que o contagem/resumo destaca.
_LISTA_PRINCIPAL = {
    "FINANCEIRO_GERAL": "resumos",
    "INDICADORES_PE": "kpis",
    "ACOMPANHAMENTO_V2": "despesas",
}

# Todas as listas de fatos do payload. A resolução de pilares varre todas —
# inclusive ``projetos``, cujos blocos trazem pilares próprios (PDI/FORMACAO/
# STARTUPS) que podem não aparecer em ``resumos``/``kpis``/``despesas``.
_LISTAS_FATOS = ("resumos", "projetos", "kpis", "despesas")

_ROTULO_METADADO = {
    "centro": "Centro de Competência",
    "periodo_referencia": "Período de referência",
    "termo": "Termo de cooperação",
}

# Ruling 22: v2 sem metadados legíveis importa com os valores do formulário,
# mas registra que a validação contra o arquivo não foi possível (aviso, nunca
# erro) — para nunca importar silenciosamente sem conferência.
_AVISO_METADADOS_AUSENTES = (
    "Metadados de centro/período não encontrados no arquivo; "
    "validação contra o arquivo não foi possível.")

# Campos por modelo: valores anuláveis (Decimal/data/int) e textos ``blank``
# (que precisam de ``""`` — o banco não aceita NULL nos CharField).
_RESUMO_VALORES = (
    "recurso_ou_meta", "captado_anos_1_2", "captado_ano3_ytd", "captado",
    "realizado", "projetado", "realizado_mais_projetado", "diferenca",
    "percentual", "percentual_meta", "percentual_captado",
)
_RESUMO_TEXTOS = ("farol",)
_PROJETO_VALORES = (
    "sequencia", "inicio", "fim", "orcado", "realizado", "projetado_2026",
    "projetado_2027", "realizado_mais_projetado", "diferenca", "percentual",
)
_PROJETO_TEXTOS = ("nome", "status")
_KPI_VALORES = (
    "sequencia", "meta_2024", "meta_2025", "meta_2026", "meta_2027", "meta_total",
    "executado_2024", "executado_2025", "acumulado", "gap",
    "projecao_2026", "projecao_2027",
)
_KPI_TEXTOS = ("nome", "descricao", "unidade")
_DESPESA_VALORES = (
    "data_nota", "data_pagamento", "data_movimento",
    "quantidade", "valor_unitario", "valor",
)
_DESPESA_TEXTOS = (
    "acao_relacionada", "codigo_projeto", "conta_projeto", "marco",
    "tipo_despesa", "credor", "documento", "numero_nota", "link_documento",
    "numero_patrimonial", "descricao", "descricao_atividade", "observacao",
    "fonte_recurso",
)


# ------------------------------------------------------------------- detecção

def detejar_tipo_arquivo(sheetnames: list[str], table_names=()) -> str | None:
    """Financeiro por objetos nomeados; demais fontes pela assinatura das abas.

    Uma tabela conhecida identifica um candidato REFAT; o parser verifica o
    contrato completo e informa precisamente quais objetos estão ausentes.
    """
    if set(table_names).intersection(financeiro_geral.TABELAS_OBRIGATORIAS):
        return "FINANCEIRO_GERAL"
    nomes = [normalizar(nome) for nome in sheetnames]
    if any(_ASSINATURA_INDICADORES in nome for nome in nomes):
        return "INDICADORES_PE"
    if any(nome in _ASSINATURA_V2 for nome in nomes):
        return "ACOMPANHAMENTO_V2"
    return None


def _mensagem_tipo(arquivo_nome: str, abas: list[str], tipo_fonte, detectado, tabelas=()) -> str:
    """Rejeição amigável citando objetos/abas exigidos por template."""
    exigidas = "; ".join(
        f"{tipo} ⇒ {' + '.join(abas_obrigatorias)}"
        for tipo, abas_obrigatorias in VALIDADORES.items())
    lidas = ", ".join(repr(a) for a in abas) if abas else "nenhuma"
    objetos = ", ".join(tabelas) if tabelas else "nenhuma"
    return (
        f'Upload rejeitado: não foi possível identificar o tipo de arquivo '
        f'"{arquivo_nome}" pela estrutura (abas lidas: {lidas}; tabelas: {objetos}). '
        f"Tipo informado: {tipo_fonte!r}; tipo detectado: {detectado!r}. "
        f"Tabelas/abas exigidas por template — {exigidas}. "
        f"Para dados financeiros, envie FINANCEIRO GERAL - REFAT.xlsx com as tabelas nomeadas."
    )


def inspecionar_arquivo(caminho: Path | str) -> tuple[list[str], list[str]]:
    """Assinatura por metadados, sem carregar as células dos outros modelos.

    O modo otimizado não expõe ws.tables. Reconstruir apenas os objetos Table
    dos metadados evita carregar integralmente as grandes abas do v2.
    """
    wb = abrir_workbook(caminho)
    try:
        with ZipFile(caminho) as pacote:
            tipos = fromstring(pacote.read("[Content_Types].xml"))
            nomes = []
            for parte in tipos:
                if parte.get("ContentType") == "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml":
                    objeto = Table.from_tree(fromstring(pacote.read(parte.get("PartName").lstrip("/"))))
                    nomes.append(objeto.displayName)
        return list(wb.sheetnames), nomes
    except Exception as exc:
        raise ArquivoInvalidoError(f"Não foi possível inspecionar '{Path(caminho).name}': {exc}") from exc
    finally:
        wb.close()


# --------------------------------------------------------------- inspeção/UI

def resumo_json(tipo_fonte: str, caminho: Path | str, centro: str, referencia: str) -> dict:
    """Só parse, sem banco: o que seria extraído, para inspeção manual."""
    if tipo_fonte not in _PARSERS:
        raise IngestaoError(_mensagem_tipo(Path(caminho).name, [], tipo_fonte, None))
    payload = _PARSERS[tipo_fonte](caminho)
    lista = payload[_LISTA_PRINCIPAL[tipo_fonte]]
    return {
        "tipo_fonte": tipo_fonte,
        "centro": centro,
        "periodo_referencia": referencia,
        "contagens": _contagens(tipo_fonte, payload),
        "exemplo_resumo": lista[0] if lista else None,
        "avisos": payload["avisos"],
        "erros": payload["erros"],
    }


# ------------------------------------------------------------------ ingestão

def ingestar(tipo_fonte: str, caminho: Path | str, centro_nome: str,
             periodo_referencia: str, usuario) -> ImportacaoAcompanhamento:
    """Detecta, valida e grava os fatos de ``tipo_fonte`` de forma atômica.

    Returns:
        ImportacaoAcompanhamento: o registro de SUCESSO da importação.

    Raises:
        IngestaoError: arquivo ilegível, tipo não identificado/divergente
            (Ruling 2) ou erros de parse/metadados — sempre após registrar o
            ``ImportacaoAcompanhamento`` de ERRO fora da transação.
        ValidationError: pilar do arquivo não existe no banco (SDD §7) — após
            rollback total e log de ERRO.
    """
    arquivo_nome = Path(caminho).name
    avisos: list[str] = []
    acomp = None
    try:
        abas, tabelas = inspecionar_arquivo(caminho)
        detectado = detejar_tipo_arquivo(abas, tabelas)
        if tipo_fonte not in _PARSERS or detectado != tipo_fonte:
            raise IngestaoError(_mensagem_tipo(arquivo_nome, abas, tipo_fonte, detectado, tabelas))
        payload = _PARSERS[tipo_fonte](caminho)
        avisos = payload["avisos"]
        if payload["erros"]:
            raise IngestaoError(" | ".join(payload["erros"]))
        _avisar_metadados_ausentes(tipo_fonte, payload)
        divergencias = _metadados_divergentes(
            payload, centro_nome, periodo_referencia, caminho)
        if divergencias:
            raise IngestaoError(" | ".join(divergencias))
        centro, _ = CentroCompetencia.objects.get_or_create(
            codigo=_codigo_centro(centro_nome), defaults={"nome": centro_nome})
        acomp, _ = Acompanhamento.objects.get_or_create(
            centro=centro, periodo_referencia=periodo_referencia)
    except IngestaoError as exc:
        _falha(acomp, tipo_fonte, arquivo_nome, usuario, avisos, str(exc))
        raise
    except ArquivoInvalidoError as exc:
        erro = IngestaoError(
            f'Upload rejeitado: não foi possível ler "{arquivo_nome}": {exc}')
        _falha(acomp, tipo_fonte, arquivo_nome, usuario, avisos, str(erro))
        raise erro from exc
    except Exception as exc:  # noqa: BLE001 — falha inesperada também vira log ERRO
        erro = IngestaoError(f"Falha ao preparar a importação de {arquivo_nome}: {exc}")
        _falha(acomp, tipo_fonte, arquivo_nome, usuario, avisos, str(erro))
        raise erro from exc

    try:
        with transaction.atomic():
            pilares = _resolver_pilares(payload)
            _gravar_fatos(acomp, tipo_fonte, payload, pilares)
            termo = (payload.get("metadados") or {}).get("termo") or ""
            if termo and acomp.termo_cooperacao != termo:
                acomp.termo_cooperacao = termo
                acomp.save()
            contagens = _contagens(tipo_fonte, payload)
            importacao = ImportacaoAcompanhamento.objects.create(
                acompanhamento=acomp,
                tipo_fonte=tipo_fonte,
                arquivo_nome=arquivo_nome,
                status="SUCESSO",
                log=_mensagem_sucesso(tipo_fonte, contagens, periodo_referencia),
                avisos=" | ".join(payload["avisos"]),
                resumo=contagens,
                usuario=usuario,
            )
            registrar_auditoria(
                usuario, "IMPORTAR_PRESTACAO", "ImportacaoAcompanhamento",
                registro_id=importacao.pk, campo=tipo_fonte,
                valor_novo=f"{contagens} · {arquivo_nome}",
            )
    except Exception as exc:  # noqa: BLE001 — todo erro vira log ERRO + rollback
        log = "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
        _falha(acomp, tipo_fonte, arquivo_nome, usuario, payload["avisos"], log)
        if isinstance(exc, (IngestaoError, ValidationError)):
            raise
        raise IngestaoError(f"Falha ao gravar os dados de {arquivo_nome}: {exc}") from exc
    return importacao


# ------------------------------------------------------------------ persistência

def _gravar_fatos(acomp, tipo_fonte: str, payload: dict, pilares: dict) -> None:
    """Apaga e recria os fatos da fonte; KPIs fazem upsert (Ruling 7)."""
    if tipo_fonte == "FINANCEIRO_GERAL":
        # Ruling 5: delete do escopo inteiro da fonte antes do insert —
        # a unicidade não protege linhas com ano=NULL.
        ResumoFinanceiro.objects.filter(acompanhamento=acomp).delete()
        ResumoFinanceiro.objects.bulk_create(
            ResumoFinanceiro(
                acompanhamento=acomp, origem=item["origem"], ano=item["ano"],
                pilar=pilares[item["pilar"]],
                **_campos(item, _RESUMO_VALORES, _RESUMO_TEXTOS))
            for item in payload["resumos"])
        ProjetoFinanceiro.objects.filter(acompanhamento=acomp).delete()
        ProjetoFinanceiro.objects.bulk_create(
            ProjetoFinanceiro(
                acompanhamento=acomp, origem=item["origem"],
                pilar=pilares[item["pilar"]],
                **_campos(item, _PROJETO_VALORES, _PROJETO_TEXTOS))
            for item in payload["projetos"])
    elif tipo_fonte == "INDICADORES_PE":
        _gravar_kpis(acomp, payload, pilares)
    else:
        DespesaAcompanhamento.objects.filter(acompanhamento=acomp).delete()
        DespesaAcompanhamento.objects.bulk_create(
            DespesaAcompanhamento(
                acompanhamento=acomp, aba=item["aba"], linha=item["linha"],
                tipo_recurso=item["tipo_recurso"], pilar=pilares[item["pilar"]],
                **_campos(item, _DESPESA_VALORES, _DESPESA_TEXTOS))
            for item in payload["despesas"])


def _gravar_kpis(acomp, payload: dict, pilares: dict) -> None:
    """Upsert por ``(acompanhamento, codigo)``: atualiza, cria e apaga o que sumiu.

    Nunca apaga e recria (Ruling 7): ``OverrideAcompanhamento.kpi`` é CASCADE e
    a recriação trocaria os PKs, destruindo em silêncio os overrides manuais
    do usuário (SDD §9 — "override > importado").
    """
    codigos = set()
    atualizar, criar = [], []
    for item in payload["kpis"]:
        codigos.add(item["codigo"])
        valores = _campos(item, _KPI_VALORES, _KPI_TEXTOS)
        valores["pilar"] = pilares[item["pilar"]]
        existente = KpiAcompanhamento.objects.filter(
            acompanhamento=acomp, codigo=item["codigo"]).first()
        if existente is None:
            criar.append(KpiAcompanhamento(
                acompanhamento=acomp, codigo=item["codigo"], **valores))
        else:
            for campo, valor in valores.items():
                setattr(existente, campo, valor)
            atualizar.append(existente)
    # Só os códigos que desapareceram do arquivo saem; os overrides dos que
    # continuam preservam PK e vínculo.
    KpiAcompanhamento.objects.filter(acompanhamento=acomp).exclude(
        codigo__in=codigos).delete()
    if atualizar:
        KpiAcompanhamento.objects.bulk_update(
            atualizar, ["pilar", *_KPI_VALORES, *_KPI_TEXTOS])
    if criar:
        KpiAcompanhamento.objects.bulk_create(criar)


def _campos(item: dict, valores: tuple[str, ...], textos: tuple[str, ...]) -> dict:
    """kwargs de modelo: campos copiados do item; textos vazios viram ``""``."""
    campos = {nome: item.get(nome) for nome in valores}
    for nome in textos:
        campos[nome] = item.get(nome) or ""
    return campos


def _resolver_pilares(payload: dict) -> dict[str, Pilar]:
    """Códigos de pilar do payload → ``Pilar``; ausente ⇒ ``ValidationError``.

    Mesma mensagem de ``apps.finance.excel_mestre.aplicar_consolidado``.
    """
    codigos = {
        item["pilar"]
        for lista in _LISTAS_FATOS
        for item in payload[lista]
    }
    pilares: dict[str, Pilar] = {}
    for codigo in sorted(codigos, key=str):
        pilar = (Pilar.objects.filter(codigo__iexact=codigo).first()
                 if codigo else None)
        if pilar is None:
            raise ValidationError(
                f"Pilar '{codigo}' não encontrado. Cadastre o pilar antes de aplicar.")
        pilares[codigo] = pilar
    return pilares


def _contagens(tipo_fonte: str, payload: dict) -> dict:
    """Contagem por tipo — o ``resumo`` gravado no log de importação."""
    if tipo_fonte == "FINANCEIRO_GERAL":
        return {"resumos": len(payload["resumos"]), "projetos": len(payload["projetos"])}
    lista = _LISTA_PRINCIPAL[tipo_fonte]
    return {lista: len(payload[lista])}


def _mensagem_sucesso(tipo_fonte: str, contagens: dict, periodo: str) -> str:
    if tipo_fonte == "FINANCEIRO_GERAL":
        return (f"{contagens['resumos']} resumo(s) e {contagens['projetos']} "
                f"projetos financeiros importados para {periodo}.")
    if tipo_fonte == "INDICADORES_PE":
        return f"{contagens['kpis']} KPIs importados para {periodo}."
    return f"{contagens['despesas']} despesas importadas para {periodo}."


def _codigo_centro(nome: str) -> str:
    """Código do centro: ``slugify`` do nome normalizado (SDD §3; Ruling 19).

    O slug real do centro tem 69 caracteres
    (``centro-de-competencia-embrapii-cimatec-em-tecnologias-quanticas-quiin``)
    e cabe no ``SlugField(max_length=200)``; nunca truncar — centros distintos
    colidiriam. ``slugify`` também garante charset válido de ``SlugField``
    para nomes com ``&``/``/``.
    """
    return slugify(normalizar(nome))


# ------------------------------------------------------------------- metadados

def _avisar_metadados_ausentes(tipo_fonte: str, payload: dict) -> None:
    """Ruling 22: v2 sem metadados legíveis importa, mas com aviso registrado.

    Só o ``ACOMPANHAMENTO_V2`` traz metadados; em ``FINANCEIRO_GERAL`` e
    ``INDICADORES_PE`` os argumentos do formulário são a fonte da verdade por
    desenho (SDD §1) e nada é avisado. Vira ``ImportacaoAcompanhamento.avisos``
    — nunca ``erros``.
    """
    if tipo_fonte != "ACOMPANHAMENTO_V2":
        return
    metadados = payload.get("metadados") or {}
    if not metadados.get("centro") or not metadados.get("periodo_referencia"):
        payload["avisos"].append(_AVISO_METADADOS_AUSENTES)


def _metadados_divergentes(payload: dict, centro_nome: str, periodo_referencia: str,
                           caminho: Path | str) -> list[str]:
    """Compara os metadados do v2 com os argumentos (SDD §1) via ``normalizar``.

    Em ``FINANCEIRO_GERAL``/``INDICADORES_PE`` os metadados são ``None``: os
    argumentos do formulário são a fonte da verdade e não há o que conferir.
    Divergência no v2 vira erro citando a aba e a célula de onde o valor foi
    lido (SDD §6/§7).
    """
    metadados = payload.get("metadados") or {}
    esperado = {"centro": centro_nome, "periodo_referencia": periodo_referencia}
    divergencias = []
    for chave, informado in esperado.items():
        lido = metadados.get(chave)
        if not lido or normalizar(lido) == normalizar(informado):
            continue
        aba, celula = _procedencia_metadados(caminho).get(chave, ("?", "?"))
        divergencias.append(
            f"Metadado divergente: {_ROTULO_METADADO[chave]} em \"{aba}\" ({celula})"
            f' = "{lido}" difere de "{informado}" informado no upload.')
    return divergencias


def _procedencia_metadados(caminho: Path | str) -> dict[str, tuple[str, str]]:
    """``chave → (aba, célula)`` de onde cada metadado do v2 foi lido.

    Reusa a varredura do próprio parser (``acompanhamento_v2._varrer_metadados``,
    com os mesmos limites) para a citação casar exatamente com o valor que o
    parser extraiu — a origem primária é o ``0. Sumário`` (SDD §6).
    """
    wb = abrir_workbook(caminho)
    try:
        origem: dict[str, tuple[str, str]] = {}
        if acompanhamento_v2.ABA_SUMARIO in wb.sheetnames:
            for chave, _valor, coordenada in acompanhamento_v2._varrer_metadados(
                    wb[acompanhamento_v2.ABA_SUMARIO],
                    acompanhamento_v2._MAX_LINHA_METADADOS_SUMARIO):
                origem.setdefault(chave, (acompanhamento_v2.ABA_SUMARIO, coordenada))
        for aba in acompanhamento_v2.ABAS_DESPESA:
            if aba not in wb.sheetnames:
                continue
            for chave, _valor, coordenada in acompanhamento_v2._varrer_metadados(
                    wb[aba], acompanhamento_v2._MAX_LINHA_METADADOS):
                origem.setdefault(chave, (aba, coordenada))
        return origem
    finally:
        wb.close()


# ------------------------------------------------------------------- falhas

def _falha(acomp, tipo_fonte, arquivo_nome: str, usuario,
           avisos: list[str], log: str) -> ImportacaoAcompanhamento:
    """Registra o ERRO **fora** da transação que falhou e audita (SDD §7).

    Ruling 20: o ``tipo_fonte`` tentado é gravado **verbatim** (as ``choices``
    são metadado de formulário, não restrição de banco) — o histórico mostra o
    que foi tentado; ``None`` (sem tipo) vira ``""`` por ser NOT NULL.
    """
    importacao = ImportacaoAcompanhamento.objects.create(
        acompanhamento=acomp,
        tipo_fonte=tipo_fonte or "",
        arquivo_nome=arquivo_nome,
        status="ERRO",
        log=log or "",
        avisos=" | ".join(avisos),
        usuario=usuario,
    )
    registrar_auditoria(
        usuario, "IMPORTAR_PRESTACAO_ERRO", "ImportacaoAcompanhamento",
        registro_id=importacao.pk, campo=tipo_fonte or "", valor_novo=log,
    )
    return importacao
