"""Testes do parser do Acompanhamento Financeiro (v2) (SDD §6).

Scaffolding em ``unittest`` (o runner do ``manage.py test`` não injeta
``tmp_path``); as asserções e valores esperados vêm literalmente do brief.
Rulings do controlador aplicados sobre o brief: Ruling 14 (``test_celula_
obrigatoria_invalida...`` planta o ``#REF!`` em coluna mapeada — a coluna
``valor`` — e não mais em Q11, não mapeada) e Ruling 15 (a regra 18 casa o
radical ``acao esta relacionad``); casos novos cobrem ambos.
"""

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from apps.accountability.parsers.acompanhamento_v2 import parse_acompanhamento_v2

RAIZ = Path(__file__).resolve().parents[3]
REAL = RAIZ / "Acompanhamento Financeiro (v2).xlsx"


def _cabecalho_metadados(ws, deslocamento=0):
    base = 5 + deslocamento
    ws.cell(base, 4, "Centro de Competência EMBRAPII:")
    ws.cell(base, 5, "Centro de Competência Embrapii CIMATEC em Tecnologias Quânticas - Quiin")
    ws.cell(base + 1, 4, "Termo de Cooperação N°:")
    ws.cell(base + 1, 5, "053/2023")
    ws.cell(base + 2, 4, "Período de Referência do Acompanhamento:")
    ws.cell(base + 2, 5, "2T/2024")


CABECALHOS = {
    "3. Conta Ação - AFCCT": (10, ["Linha", "Informação do código do projeto de AFCCT",
                                   "Conta do projeto", "Tipo de despesa", "Credor",
                                   "Data do pagamento", "Valor (R$)"],
                              [1, "PDI-01", "AFCCT-01", "Serviço", "Fornecedor X", "15/04/2024", 123.45]),
    "4. Conta Ação - FCRH": (9, ["Linha", "Informação do código do projeto de FCRH",
                                 "Tipo de despesa", "Credor", "Número da Nota Fiscal ou Invoice",
                                 "Data da Nota Fiscal ou Invoice", "Data do pagamento", "Valor (R$)"],
                             [1, "FCRH-01", "Bolsa", "Fornecedor Y", "NF-100", "01/04/2024",
                              "15/04/2024", 123.45]),
    "5. Conta Ação - ACS": (9, ["Linha", "Informação do código do projeto de ACS",
                                "Marco do projeto, se aplicável", "Credor",
                                "Data do pagamento", "Valor (R$)"],
                            [1, "ACS-01", "M1", "Fornecedor Z", "15/04/2024", 123.45]),
    "6. Conta Ação - AT": (10, ["Linha", "Informar a qual ação está relacionada",
                                "Informar o código do projeto/atividade", "Tipo de despesa",
                                "Entidade que realizou o aporte", "Descrição", "Data", "Valor (R$)"],
                           [1, "Associação", "AT-01", "Aporte", "Empresa A", "Aporte", "15/04/2024", 123.45]),
    "6.1 Conta Ação - AT (Lei TICs)": (10, ["Linha", "Informar a qual ação está relacionada",
                                            "Informar o código do projeto/atividade", "Tipo de despesa",
                                            "Entidade que realizou o aporte", "Descrição",
                                            "Data", "Valor (R$)"],
                                       [1, "Associação", "AT-02", "Aporte", "Empresa B", "Aporte",
                                        "15/04/2024", 123.45]),
    "7. Conta Ação - Outras Fontes": (10, ["Linha", "Informar a qual ação está relacionada",
                                           "Informar o código do projeto/atividade", "Tipo de despesa",
                                           "Entidade que realizou o aporte", "Descrição",
                                           "Data", "Valor (R$)"],
                                      [1, "Outra fonte", "OF-01", "Aporte", "Empresa C", "Aporte",
                                       "15/04/2024", 123.45]),
    "8. Conta Ação - Infraestrutura": (10, ["Linha", "Tipo de despesa", "Fornecedor – estrangeiro?",
                                            "Nome (Fornecedor)", "CNPJ (Fornecedor)",
                                            "Número da Nota ou Invoice", "Data da Nota Fiscal ou Invoice",
                                            "Data do pagamento",
                                            "Valor da Nota Fiscal ou Invoice (RS)", "Observações"],
                                       [1, "Equipamento", "Não", "Fornecedor D", "11.222.333/0001-44",
                                        "NF-200", "01/04/2024", "15/04/2024", 123.45, "ok"]),
    "9. Ampliação de Infraestrutura": (9, ["Linha", "Descrição do item", "Categoria do item",
                                          "Fonte Recurso", "Quantidade",
                                          "Valor unitário do item na Nota Fiscal ou Invoice (R$)",
                                          "Valor total dos itens na Nota Fiscal ou Invoice (R$)",
                                          "Número patrimonial do bem", "Observações"],
                                      [1, "Servidor", "TI", "EMBRAPII", 2, 61.725, 123.45, "PAT-9", "ok"]),
}


def _workbook_despesas():
    wb = Workbook()
    wb.remove(wb.active)
    ws0 = wb.create_sheet("0. Sumário")
    _cabecalho_metadados(ws0, deslocamento=14)
    for nome, (linha_cab, cabecalhos, dados) in CABECALHOS.items():
        aba = wb.create_sheet(nome)
        _cabecalho_metadados(aba)
        for i, texto in enumerate(cabecalhos, start=1):
            aba.cell(linha_cab, i, texto)
        for i, valor in enumerate(dados, start=1):
            aba.cell(linha_cab + 1, i, valor)
        aba.cell(linha_cab + 2, 1, 2)   # só o número sequencial: encerra a leitura
        aba.cell(40, 15, "#REF!")       # coluna não mapeada: vira aviso, nunca erro
    return wb


class ParseAcompanhamentoV2Testes(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_metadados_do_arquivo_real(self):
        p = parse_acompanhamento_v2(REAL)
        assert p["erros"] == []
        assert p["metadados"]["centro"].startswith("Centro de Competência Embrapii CIMATEC")
        assert p["metadados"]["termo"] == "053/2023"
        assert p["metadados"]["periodo_referencia"] == "2T/2024"

    def test_arquivo_real_sem_lancamentos_nao_gera_erro(self):
        p = parse_acompanhamento_v2(REAL)
        assert p["despesas"] == []
        assert p["erros"] == []

    def test_fixture_gera_uma_despesa_por_aba(self):
        caminho = Path(self.tmp.name) / "v2.xlsx"
        _workbook_despesas().save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"] == []
        assert len(p["despesas"]) == 8
        por_aba = {d["aba"]: d for d in p["despesas"]}
        assert por_aba["3. Conta Ação - AFCCT"]["pilar"] == "PDI"
        assert por_aba["3. Conta Ação - AFCCT"]["codigo_projeto"] == "PDI-01"
        assert por_aba["4. Conta Ação - FCRH"]["numero_nota"] == "NF-100"
        assert por_aba["4. Conta Ação - FCRH"]["data_nota"].isoformat() == "2024-04-01"
        assert por_aba["5. Conta Ação - ACS"]["pilar"] == "STARTUPS"
        assert por_aba["7. Conta Ação - Outras Fontes"]["pilar"] == "OUTRASFONTES"
        assert por_aba["8. Conta Ação - Infraestrutura"]["documento"] == "11.222.333/0001-44"
        assert por_aba["8. Conta Ação - Infraestrutura"]["credor"] == "Fornecedor D"
        assert por_aba["9. Ampliação de Infraestrutura"]["valor"] == Decimal("123.45")
        assert por_aba["9. Ampliação de Infraestrutura"]["valor_unitario"] == Decimal("61.725")
        assert por_aba["9. Ampliação de Infraestrutura"]["quantidade"] == Decimal("2")
        assert por_aba["9. Ampliação de Infraestrutura"]["numero_patrimonial"] == "PAT-9"
        assert all(d["valor"] == Decimal("123.45") for d in p["despesas"])
        assert not [d for d in p["despesas"] if d["linha"] > 11]

    def test_lei_tics_segregada_do_at(self):
        caminho = Path(self.tmp.name) / "v2.xlsx"
        _workbook_despesas().save(caminho)
        p = parse_acompanhamento_v2(caminho)
        tipos = {d["aba"]: d["tipo_recurso"] for d in p["despesas"]}
        assert tipos["6. Conta Ação - AT"] == "AT"
        assert tipos["6.1 Conta Ação - AT (Lei TICs)"] == "AT_LEI_TICS"

    def test_metadados_divergentes_entre_abas_rejeitam(self):
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        wb["3. Conta Ação - AFCCT"].cell(7, 5, "3T/2024")
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"], "divergência de metadado deve virar erro"
        assert any("AFCCT" in e for e in p["erros"])

    def test_aba_obrigatoria_ausente_gera_erro_amigavel(self):
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        wb.remove(wb["6.1 Conta Ação - AT (Lei TICs)"])
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert any("6.1 Conta Ação - AT (Lei TICs)" in e for e in p["erros"])

    def test_celula_obrigatoria_invalida_aponta_linha_e_coluna(self):
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        # Ruling 14: o erro tem de estar em coluna MAPEADA — a coluna `valor`
        # (coluna G) da linha de dados da aba 3 na fixture.
        wb["3. Conta Ação - AFCCT"].cell(11, 7, "#REF!")
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert any("AFCCT" in e and "linha" in e.lower() for e in p["erros"])

    def test_ref_em_coluna_nao_mapeada_de_linha_de_dados_vira_aviso(self):
        # Ruling 14: o template real traz "#REF!" na coluna "Ano" (não mapeada)
        # da aba 8 em todas as linhas; não pode bloquear a importação (SDD §7).
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        aba = wb["8. Conta Ação - Infraestrutura"]
        aba.cell(10, 15, "Ano")
        aba.cell(11, 15, "#REF!")
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"] == []
        assert any("#REF!" in a for a in p["avisos"])
        assert len(p["despesas"]) == 8

    def test_cabecalho_real_acao_relacionado_masculino_mapeia(self):
        # Ruling 15: o template real escreve "Informar a qual ação está
        # relacionado" (masculino); a regra 18 casa o radical "acao esta
        # relacionad" e cobre os dois gêneros.
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        wb["6.1 Conta Ação - AT (Lei TICs)"].cell(10, 2, "Informar a qual ação está relacionado")
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"] == []
        por_aba = {d["aba"]: d for d in p["despesas"]}
        assert por_aba["6.1 Conta Ação - AT (Lei TICs)"]["acao_relacionada"] == "Associação"
        assert por_aba["6. Conta Ação - AT"]["acao_relacionada"] == "Associação"

    def test_quantidade_e_valor_unitario_preservam_precisao(self):
        # Guarda da Ruling 11: quantidade/valor_unitário não são moeda —
        # 61.725 (3 casas) e quantidades fracionárias não podem ser quantizados.
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        wb["9. Ampliação de Infraestrutura"].cell(10, 5, 2.5)
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"] == []
        ampliacao = next(d for d in p["despesas"]
                         if d["aba"] == "9. Ampliação de Infraestrutura")
        assert ampliacao["valor_unitario"] == Decimal("61.725")
        assert ampliacao["quantidade"] == Decimal("2.5")
        assert ampliacao["valor"] == Decimal("123.45")

    def test_conteudo_abaixo_do_terminador_vira_aviso(self):
        # Item 1: a linha com só a sequência encerra a leitura (regra mantida),
        # mas conteúdo real abaixo dela não pode sumir em silêncio.
        caminho = Path(self.tmp.name) / "v2.xlsx"
        wb = _workbook_despesas()
        wb["3. Conta Ação - AFCCT"].cell(13, 7, 99.99)  # abaixo do terminador (linha 12)
        wb.save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"] == []
        assert len(p["despesas"]) == 8
        assert not [d for d in p["despesas"]
                    if d["aba"] == "3. Conta Ação - AFCCT" and d["linha"] == 13]
        assert any("3. Conta Ação - AFCCT" in a and "encerrou na linha 12" in a
                   for a in p["avisos"])

    def test_conta_projeto_capturada_ao_lado_do_codigo(self):
        # Ruling 17: a regra 2 foi dividida — "Conta do projeto" vira
        # `conta_projeto` em vez de ser descartada por "primeira coluna vence".
        caminho = Path(self.tmp.name) / "v2.xlsx"
        _workbook_despesas().save(caminho)
        p = parse_acompanhamento_v2(caminho)
        assert p["erros"] == []
        afcct = next(d for d in p["despesas"] if d["aba"] == "3. Conta Ação - AFCCT")
        assert afcct["codigo_projeto"] == "PDI-01"
        assert afcct["conta_projeto"] == "AFCCT-01"


if __name__ == "__main__":
    unittest.main()
