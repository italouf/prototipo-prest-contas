"""Fixtures REFAT com tabelas reais e valores já calculados pelo Excel."""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

FINANCEIRO = (Path(__file__).resolve().parents[3] / "mockup" / "exemplos_arquivos"
              / "FINANCEIRO GERAL - REFAT.xlsx")


def workbook_financeiro(vazio=False):
    # data_only preserva os números salvos em vez de gerar fórmulas sem cache.
    wb = load_workbook(FINANCEIRO, data_only=True)
    if vazio:
        for ws in wb:
            for tabela in ws.tables.values():
                primeira, inicio, ultima, fim = range_boundaries(tabela.ref)
                for linha in ws.iter_rows(min_col=primeira, max_col=ultima,
                                          min_row=inicio + 1, max_row=fim):
                    for celula in linha:
                        celula.value = None
    return wb


def tabela_financeira(wb, nome):
    return next((ws, ws.tables[nome]) for ws in wb if nome in ws.tables)


def preencher_tabela(wb, nome, registros):
    """Escreve registros com cabeçalhos literais e redimensiona o objeto Table."""
    ws, tabela = tabela_financeira(wb, nome)
    primeira, inicio, ultima, fim = range_boundaries(tabela.ref)
    cabecalhos = [coluna.name for coluna in tabela.tableColumns]
    for linha in ws.iter_rows(min_col=primeira, max_col=ultima,
                              min_row=inicio + 1, max_row=fim):
        for celula in linha:
            celula.value = None
    for indice, registro in enumerate(registros, start=inicio + 1):
        for coluna, cabecalho in enumerate(cabecalhos, start=primeira):
            ws.cell(indice, coluna).value = registro.get(cabecalho)
    tabela.ref = (f"{get_column_letter(primeira)}{inicio}:"
                  f"{get_column_letter(ultima)}{max(inicio, inicio + len(registros))}")
    return ws, tabela
