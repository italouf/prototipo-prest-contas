import json
from decimal import Decimal
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase

from apps.core.permissions import adicionar_grupo, garantir_grupos
from apps.pillars.models import Pilar
from .. import services
from ..models import (Acompanhamento, ImportacaoAcompanhamento, KpiAcompanhamento,
                     ProjetoFinanceiro, ResumoFinanceiro)

RAIZ = Path(__file__).resolve().parents[3]
FINANCEIRO = RAIZ / "mockup" / "exemplos_arquivos" / "FINANCEIRO GERAL - REFAT.xlsx"
INDICADORES = RAIZ / "mockup" / "exemplos_arquivos" / "Indicadores Gerais do Termo de Retificação do PE.xlsx"
V2 = RAIZ / "Acompanhamento Financeiro (v2).xlsx"

CENTRO = "Centro de Competência Embrapii CIMATEC em Tecnologias Quânticas - Quiin"
REFERENCIA = "2T/2024"


class DeteccaoTipoTestes(TestCase):
    def test_identifica_os_tres_tipos_pelas_abas(self):
        assert services.detejar_tipo_arquivo(["HEAD - INDICADORES - EMBRAPII"]) == "INDICADORES_PE"
        assert services.detejar_tipo_arquivo(["db_geral"], ["tbl_ConsolidadoPrograma"]) == "FINANCEIRO_GERAL"
        assert services.detejar_tipo_arquivo(["FINANCEIRO "]) is None
        assert services.detejar_tipo_arquivo(
            ["0. Sumário", "3. Conta Ação - AFCCT"]) == "ACOMPANHAMENTO_V2"
        assert services.detejar_tipo_arquivo(["Qualquer coisa"]) is None

    def test_arquivo_nao_identificado_vira_erro_amigavel(self):
        with self.assertRaises(services.IngestaoError) as ctx:
            services.ingestar("inventado", V2, CENTRO, REFERENCIA, None)
        self.assertIn("não foi possível identificar", str(ctx.exception))


class IngestaoFinanceiroTestes(TestCase):
    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        User = get_user_model()
        cls.user = adicionar_grupo(
            User.objects.create_user(username="ingest", password="x"), "Master")
        for codigo, nome in (("PDI", "PDI"), ("FORMACAO", "Formação FCRH"),
                             ("STARTUPS", "ACS"), ("INFRA", "Infraestrutura"),
                             ("AT", "AT"), ("OUTRASFONTES", "Outras Fontes")):
            Pilar.objects.create(codigo=codigo, nome=nome)

    def test_extrai_e_grava_os_valores_esperados(self):
        imp = services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(imp.status, "SUCESSO")
        self.assertEqual(imp.resumo, {"resumos": 13, "projetos": 33})
        pdi = ResumoFinanceiro.objects.get(origem="TAB. 1", pilar__codigo="PDI", ano=None)
        self.assertEqual(pdi.recurso_ou_meta, Decimal("29000000"))
        self.assertEqual(pdi.realizado, Decimal("17675109.56"))
        self.assertEqual(ProjetoFinanceiro.objects.filter(origem="TAB. 2").count(), 18)
        self.assertEqual(ProjetoFinanceiro.objects.filter(origem="TAB. 6").count(), 7)
        at = ResumoFinanceiro.objects.get(origem="TAB. 9", pilar__codigo="AT", ano=None)
        self.assertEqual(at.captado, Decimal("5616700"))
        self.assertNotIn("erro", imp.status.lower())

    def test_reenvio_do_mesmo_arquivo_nao_duplica(self):
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(ResumoFinanceiro.objects.count(), 13)
        self.assertEqual(ProjetoFinanceiro.objects.count(), 33)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="SUCESSO").count(), 2)

    def test_reimportar_uma_fonte_preserva_as_demas(self):
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        services.ingestar("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA, self.user)
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(ResumoFinanceiro.objects.count(), 13)
        self.assertEqual(KpiAcompanhamento.objects.count(), 10)

    def test_v2_real_grava_acompanhamento_com_metadados(self):
        services.ingestar("ACOMPANHAMENTO_V2", V2, CENTRO, REFERENCIA, self.user)
        acomp = Acompanhamento.objects.get()
        self.assertEqual(acomp.periodo_referencia, REFERENCIA)
        self.assertEqual(acomp.termo_cooperacao, "053/2023")
        self.assertEqual(acomp.despesas.count(), 0)

    def test_imprime_resumo_json_do_que_foi_extraido(self):
        payload = services.resumo_json("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA)
        texto = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
        assert payload["tipo_fonte"] == "FINANCEIRO_GERAL"
        assert payload["periodo_referencia"] == REFERENCIA
        assert payload["contagens"] == {"resumos": 13, "projetos": 33}
        assert str(payload["exemplo_resumo"]["realizado"]) == "17675109.56"
        assert '"realizado": "17675109.56"' in texto

    def test_falha_na_gravacao_faz_rollback_total(self):
        Pilar.objects.all().delete()
        with self.assertRaises(ValidationError):
            services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(ResumoFinanceiro.objects.count(), 0)
        self.assertEqual(ProjetoFinanceiro.objects.count(), 0)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 1)
        self.assertEqual(Acompanhamento.objects.count(), 1)  # histórico de erro é preservado

    def test_incluir_e_remover_projeto_atualiza_contagem_sem_duplicar(self):
        from openpyxl.utils.cell import range_boundaries
        from .fixtures_financeiro import workbook_financeiro, tabela_financeira, preencher_tabela
        from .. import panel

        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        primeira, inicio, ultima, fim = range_boundaries(tabela.ref)
        cabecalhos = [c.name for c in tabela.tableColumns]
        registros = [dict(zip(cabecalhos, linha)) for linha in ws.iter_rows(
            min_col=primeira, max_col=ultima, min_row=inicio + 1, max_row=fim, values_only=True)]
        registros.append({"PROJETO": "Projeto Novo", "ORÇADO SÍNTESE": 100})
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "outro-nome.xlsx"
            preencher_tabela(wb, "tbl_ProjetosPDI", registros)
            wb.save(caminho)
            for _ in range(2):
                imp = services.ingestar("FINANCEIRO_GERAL", caminho, CENTRO, REFERENCIA, self.user)
                self.assertEqual(imp.resumo, {"resumos": 13, "projetos": 34})
            acomp = Acompanhamento.objects.get()
            pdi = Pilar.objects.get(codigo="PDI")
            self.assertEqual(panel.painel_pilar_acompanhamento(pdi, None, "fin", acomp)["projetos"], 19)
            self.client.force_login(self.user)
            resposta = self.client.get(f"/pilar/{pdi.pk}/?ano=todos&base=fin&acompanhamento={acomp.pk}")
            self.assertContains(resposta, 'data-testid="pilar-projetos"')
            self.assertContains(resposta, "19 <span")
            preencher_tabela(wb, "tbl_ProjetosPDI", registros[:-1])
            wb.save(caminho)
            imp = services.ingestar("FINANCEIRO_GERAL", caminho, CENTRO, REFERENCIA, self.user)
            self.assertEqual(imp.resumo["projetos"], 33)
            self.assertFalse(ProjetoFinanceiro.objects.filter(nome="Projeto Novo").exists())
        wb.close()

    def test_formato_antigo_e_estrutura_incompleta_preservam_dados(self):
        from .fixtures_financeiro import workbook_financeiro, tabela_financeira
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        antes = _instantaneo()
        with self.assertRaises(services.IngestaoError) as ctx:
            services.ingestar("FINANCEIRO_GERAL", FINANCEIRO.with_name("FINANCEIRO GERAL.xlsx"), CENTRO, REFERENCIA, self.user)
        self.assertIn("REFAT", str(ctx.exception))
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        del ws.tables[tabela.displayName]
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "incompleto.xlsx"
            wb.save(caminho)
            with self.assertRaises(services.IngestaoError) as ctx:
                services.ingestar("FINANCEIRO_GERAL", caminho, CENTRO, REFERENCIA, self.user)
            self.assertIn("tbl_ProjetosPDI", str(ctx.exception))
        wb.close()
        self.assertEqual(_instantaneo(), antes)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 2)


# ---------------------------------------------------------------------------
# Adições ao piso do brief (Rulings 2, 5 e 7, metadados e auditoria). Os testes
# acima são mantidos literalmente como no brief; estes complementam — nunca
# substituem. Os imports extras ficam aqui para não tocar no bloco do brief.
import tempfile  # noqa: E402
from datetime import date  # noqa: E402

from openpyxl import Workbook  # noqa: E402

from apps.audit.models import AuditLog  # noqa: E402

from ..models import (CentroCompetencia, DespesaAcompanhamento,  # noqa: E402
                      OverrideAcompanhamento)


class RulingsIngestaoTestes(TestCase):
    """Comportamentos exigidos pelos rulings que o brief não cobre."""

    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        User = get_user_model()
        cls.user = adicionar_grupo(
            User.objects.create_user(username="rulings", password="x"), "Master")
        for codigo, nome in (("PDI", "PDI"), ("FORMACAO", "Formação FCRH"),
                             ("STARTUPS", "ACS"), ("INFRA", "Infraestrutura"),
                             ("AT", "AT"), ("OUTRASFONTES", "Outras Fontes")):
            Pilar.objects.create(codigo=codigo, nome=nome)

    def test_tipo_divergente_do_detectado_rejeita_com_mensagem_amigavel(self):
        with self.assertRaises(services.IngestaoError) as ctx:
            services.ingestar("INDICADORES_PE", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertIn("não foi possível identificar", str(ctx.exception))
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 1)

    def test_tipo_none_do_upload_rejeita_com_a_mesma_mensagem(self):
        with self.assertRaises(services.IngestaoError) as ctx:
            services.ingestar(None, V2, CENTRO, REFERENCIA, self.user)
        self.assertIn("não foi possível identificar", str(ctx.exception))

    def test_rejeicao_registra_o_tipo_tentado_verbatim(self):
        with self.assertRaises(services.IngestaoError):
            services.ingestar("inventado", V2, CENTRO, REFERENCIA, self.user)
        log = ImportacaoAcompanhamento.objects.get(status="ERRO")
        self.assertEqual(log.tipo_fonte, "inventado")  # Ruling 20: sem coerção

    def test_ingestao_atribui_o_slug_do_centro_real(self):
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(
            Acompanhamento.objects.get().centro.codigo,
            "centro-de-competencia-embrapii-cimatec-em-tecnologias-quanticas-quiin")

    def test_kpi_upsert_preserva_pk_e_overrides_do_usuario(self):
        services.ingestar("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA, self.user)
        kpi = KpiAcompanhamento.objects.get(codigo="PE-01")
        acomp = Acompanhamento.objects.get()
        OverrideAcompanhamento.objects.create(
            acompanhamento=acomp, base="fisico", kpi=kpi, ano=2025, campo="previsto",
            valor=Decimal("7"), chave="fis|PE-01|2025|previsto", usuario=self.user)
        services.ingestar("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA, self.user)
        self.assertEqual(KpiAcompanhamento.objects.get(codigo="PE-01").pk, kpi.pk)
        # recriar o KPI dispararia o CASCADE do override (Ruling 7)
        self.assertEqual(OverrideAcompanhamento.objects.get().kpi_id, kpi.pk)
        self.assertEqual(KpiAcompanhamento.objects.count(), 10)

    def test_kpi_que_sumiu_do_arquivo_e_apagado(self):
        services.ingestar("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA, self.user)
        KpiAcompanhamento.objects.create(
            acompanhamento=Acompanhamento.objects.get(), codigo="PE-99", sequencia=99,
            nome="Extinto", pilar=Pilar.objects.get(codigo="PDI"))
        services.ingestar("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA, self.user)
        self.assertFalse(KpiAcompanhamento.objects.filter(codigo="PE-99").exists())
        self.assertEqual(KpiAcompanhamento.objects.count(), 10)

    def test_metadado_divergente_rejeita_citando_aba_e_celula(self):
        with self.assertRaises(services.IngestaoError) as ctx:
            services.ingestar("ACOMPANHAMENTO_V2", V2, CENTRO, "1T/2024", self.user)
        mensagem = str(ctx.exception)
        self.assertIn("0. Sumário", mensagem)
        self.assertIn("2T/2024", mensagem)
        self.assertIn("1T/2024", mensagem)
        self.assertRegex(mensagem, r"\([A-Z]+\d+\)")  # cita a célula de origem
        self.assertEqual(Acompanhamento.objects.count(), 0)  # nada persiste
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 1)

    def test_auditoria_registrada_no_sucesso_e_no_erro(self):
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(AuditLog.objects.filter(acao="IMPORTAR_PRESTACAO").count(), 1)
        Pilar.objects.all().delete()
        with self.assertRaises(ValidationError):
            services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        self.assertEqual(AuditLog.objects.filter(acao="IMPORTAR_PRESTACAO_ERRO").count(), 1)

    def test_resumo_json_dos_indicadores(self):
        payload = services.resumo_json("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA)
        self.assertEqual(payload["tipo_fonte"], "INDICADORES_PE")
        self.assertEqual(payload["contagens"], {"kpis": 10})


class CentroCodigoTestes(TestCase):
    def test_codigo_do_centro_real_cabe_no_max_length(self):
        from django.core.exceptions import ValidationError
        centro = CentroCompetencia.objects.create(
            codigo="centro-de-competencia-embrapii-cimatec-em-tecnologias-quanticas-quiin",
            nome="Centro de Competência Embrapii CIMATEC em Tecnologias Quânticas - Quiin")
        centro.full_clean()   # levanta se codigo exceder max_length
        self.assertLessEqual(
            len(centro.codigo), CentroCompetencia._meta.get_field("codigo").max_length)


# ------------------------------------------------------- fixtures sintéticas

def _workbook_financeiro(caminho, nomes_projetos):
    """REFAT em que PDI só aparece nos projetos, para testar resolução de pilares."""
    from .fixtures_financeiro import workbook_financeiro, preencher_tabela

    wb = workbook_financeiro(vazio=True)
    preencher_tabela(wb, "tbl_ConsolidadoPrograma", [{
        "AÇÃO": "INFRAESTRUTURA", "RECURSO PPI TOTAL": 100, "REALIZADO TOTAL": 50,
    }])
    preencher_tabela(wb, "tbl_ProjetosPDI", [
        {"PROJETO": nome, "STATUS": "Em execução", "ORÇADO SÍNTESE": 10,
         "REALIZADO (até 05/2026)": 5} for nome in nomes_projetos
    ])
    wb.save(caminho)
    wb.close()


# Abas obrigatórias do v2 (SDD §6): sem as 8 o parser acusa aba ausente.
_ABAS_V2 = (
    "3. Conta Ação - AFCCT",
    "4. Conta Ação - FCRH",
    "5. Conta Ação - ACS",
    "6. Conta Ação - AT",
    "6.1 Conta Ação - AT (Lei TICs)",
    "7. Conta Ação - Outras Fontes",
    "8. Conta Ação - Infraestrutura",
    "9. Ampliação de Infraestrutura",
)

# Uma linha de despesa nas abas-alvo, com as colunas que o SDD §6 mapeia
# (inclui quantidade/valor unitário de 3 casas e data para o teste de round trip).
_DADOS_V2 = {
    "3. Conta Ação - AFCCT": (
        ["Linha", "Informação do código do projeto de AFCCT", "Conta do projeto",
         "Data do pagamento", "Credor", "Valor (R$)", "Observação"],
        [1, "PDI-01", "CTA-1", date(2024, 4, 15), "Fornecedor X", 123.45, "ok"]),
    "9. Ampliação de Infraestrutura": (
        ["Linha", "Descrição do item", "Quantidade",
         "Valor unitário do item na Nota Fiscal ou Invoice (R$)",
         "Valor total dos itens na Nota Fiscal ou Invoice (R$)",
         "Número patrimonial do bem"],
        [1, "Servidor", 2, 61.725, 123.45, "PAT-9"]),
}


def _workbook_v2(caminho, com_metadados=True):
    """Workbook do Acompanhamento Financeiro (v2) com despesas nas abas 3 e 9.

    Estilo das fixtures de ``test_parsers_acompanhamento``: as outras abas
    trazem só o cabeçalho ``Linha`` (aba e cabeçalho são obrigatórios) e os
    rótulos de metadados ficam no ``0. Sumário``, fonte primária da conferência
    de centro/período (SDD §6). ``com_metadados=False`` simula um template com
    rótulos ilegíveis (Ruling 22).
    """
    wb = Workbook()
    wb.remove(wb.active)
    ws0 = wb.create_sheet("0. Sumário")
    if com_metadados:
        ws0.cell(2, 2, "Centro de Competência")
        ws0.cell(2, 3, CENTRO)
        ws0.cell(3, 2, "Período de referência")
        ws0.cell(3, 3, REFERENCIA)
        ws0.cell(4, 2, "Termo de cooperação")
        ws0.cell(4, 3, "053/2023")
    for nome in _ABAS_V2:
        aba = wb.create_sheet(nome)
        cabecalhos, dados = _DADOS_V2.get(nome, (["Linha"], [1]))
        for coluna, texto in enumerate(cabecalhos, start=1):
            aba.cell(5, coluna, texto)
        for coluna, valor in enumerate(dados, start=1):
            aba.cell(6, coluna, valor)
    wb.save(caminho)


def _instantaneo():
    """Snapshot das linhas de resumo/projeto para comparar antes e depois."""
    resumos = list(ResumoFinanceiro.objects.order_by(
        "origem", "pilar__codigo", "ano").values_list(
            "origem", "pilar__codigo", "ano", "recurso_ou_meta", "captado",
            "realizado"))
    projetos = list(ProjetoFinanceiro.objects.order_by(
        "origem", "sequencia", "nome").values_list(
            "origem", "pilar__codigo", "sequencia", "nome", "orcado",
            "realizado"))
    return resumos, projetos


class PilarSoEmProjetosTestes(TestCase):
    """Pilar que só aparece em ``projetos`` precisa entrar na resolução."""

    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        User = get_user_model()
        cls.user = adicionar_grupo(
            User.objects.create_user(username="pilarprojeto", password="x"), "Master")
        Pilar.objects.create(codigo="INFRA", nome="Infraestrutura")

    def test_pilar_somente_de_projeto_e_resolvido(self):
        Pilar.objects.create(codigo="PDI", nome="PDI")
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "sintetico.xlsx"
            _workbook_financeiro(caminho, ["Projeto Alfa"])
            imp = services.ingestar("FINANCEIRO_GERAL", caminho, CENTRO, REFERENCIA, self.user)
        self.assertEqual(imp.status, "SUCESSO")
        self.assertEqual(imp.resumo, {"resumos": 1, "projetos": 1})
        self.assertEqual(ProjetoFinanceiro.objects.get().pilar.codigo, "PDI")
        self.assertEqual(ResumoFinanceiro.objects.get().pilar.codigo, "INFRA")

    def test_pilar_somente_de_projeto_ausente_vira_validation_error(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "sintetico.xlsx"
            _workbook_financeiro(caminho, ["Projeto Alfa"])
            with self.assertRaises(ValidationError) as ctx:
                services.ingestar("FINANCEIRO_GERAL", caminho, CENTRO, REFERENCIA, self.user)
        self.assertIn(  # mensagem mandatória, não KeyError
            "Pilar 'PDI' não encontrado. Cadastre o pilar antes de aplicar.",
            str(ctx.exception))
        self.assertEqual(ResumoFinanceiro.objects.count(), 0)
        self.assertEqual(ProjetoFinanceiro.objects.count(), 0)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 1)


class RollbackNoMeioDaGravacaoTestes(TestCase):
    """Falha depois de gravar fatos dentro da transação não pode deixar parcial."""

    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        User = get_user_model()
        cls.user = adicionar_grupo(
            User.objects.create_user(username="rollbackreal", password="x"), "Master")
        for codigo, nome in (("PDI", "PDI"), ("FORMACAO", "Formação FCRH"),
                             ("STARTUPS", "ACS"), ("INFRA", "Infraestrutura"),
                             ("AT", "AT"), ("OUTRASFONTES", "Outras Fontes")):
            Pilar.objects.create(codigo=codigo, nome=nome)

    def test_falha_apos_gravar_fatos_restaura_as_linhas_anteriores(self):
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        resumos_antes, projetos_antes = _instantaneo()
        self.assertEqual(len(resumos_antes), 13)
        self.assertEqual(len(projetos_antes), 33)

        # Simula falha no bulk_create DEPOIS de recriar os resumos; validação
        # de projetos duplicados agora acontece antes da transação.
        from unittest.mock import patch
        from django.db import IntegrityError
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "duplicado.xlsx"
            _workbook_financeiro(caminho, ["Projeto Beta"])
            with patch.object(ProjetoFinanceiro.objects, "bulk_create", side_effect=IntegrityError("falha simulada")), self.assertRaises(services.IngestaoError):
                services.ingestar("FINANCEIRO_GERAL", caminho, CENTRO, REFERENCIA, self.user)

        resumos_depois, projetos_depois = _instantaneo()
        self.assertEqual(resumos_depois, resumos_antes)   # linhas pré-existentes idênticas
        self.assertEqual(projetos_depois, projetos_antes)
        self.assertFalse(ProjetoFinanceiro.objects.filter(nome="Projeto Beta").exists())
        self.assertEqual(ResumoFinanceiro.objects.count(), 13)
        self.assertEqual(ProjetoFinanceiro.objects.count(), 33)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="SUCESSO").count(), 1)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 1)


class DespesasV2Testes(TestCase):
    """Parser → services → banco para ``DespesaAcompanhamento``.

    O v2 real não tem lançamentos (SDD §11): sem esta fixture o caminho de
    gravação de despesas — a razão de existir da metade v2 — nunca seria
    exercitado de ponta a ponta.
    """

    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        User = get_user_model()
        cls.user = adicionar_grupo(
            User.objects.create_user(username="despesasv2", password="x"), "Master")
        for codigo, nome in (("PDI", "PDI"), ("FORMACAO", "Formação FCRH"),
                             ("STARTUPS", "ACS"), ("INFRA", "Infraestrutura"),
                             ("AT", "AT"), ("OUTRASFONTES", "Outras Fontes")):
            Pilar.objects.create(codigo=codigo, nome=nome)

    def test_despesas_gravadas_com_valores_e_precisao(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "v2.xlsx"
            _workbook_v2(caminho)
            imp = services.ingestar("ACOMPANHAMENTO_V2", caminho, CENTRO, REFERENCIA, self.user)
        self.assertEqual(imp.status, "SUCESSO")
        self.assertEqual(imp.resumo, {"despesas": 2})
        self.assertNotIn("não encontrados no arquivo", imp.avisos)  # metadados conferidos
        self.assertEqual(DespesaAcompanhamento.objects.count(), 2)

        afcct = DespesaAcompanhamento.objects.get(aba="3. Conta Ação - AFCCT")
        self.assertEqual(afcct.linha, 6)
        self.assertEqual(afcct.pilar.codigo, "PDI")
        self.assertEqual(afcct.tipo_recurso, "EMBRAPII")
        self.assertEqual(afcct.codigo_projeto, "PDI-01")
        self.assertEqual(afcct.conta_projeto, "CTA-1")
        self.assertEqual(afcct.credor, "Fornecedor X")
        self.assertEqual(afcct.data_pagamento, date(2024, 4, 15))
        self.assertEqual(afcct.valor, Decimal("123.45"))
        self.assertEqual(afcct.observacao, "ok")

        ampliacao = DespesaAcompanhamento.objects.get(aba="9. Ampliação de Infraestrutura")
        self.assertEqual(ampliacao.linha, 6)
        self.assertEqual(ampliacao.pilar.codigo, "INFRA")
        self.assertEqual(ampliacao.tipo_recurso, "EMBRAPII")
        self.assertEqual(ampliacao.descricao, "Servidor")
        self.assertEqual(ampliacao.quantidade, Decimal("2"))
        self.assertEqual(ampliacao.valor_unitario, Decimal("61.725"))  # 3 casas preservadas
        self.assertEqual(ampliacao.valor, Decimal("123.45"))
        self.assertEqual(ampliacao.numero_patrimonial, "PAT-9")

    def test_reimportar_v2_substitui_despesas_sem_tocar_nas_demas_fontes(self):
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, CENTRO, REFERENCIA, self.user)
        services.ingestar("INDICADORES_PE", INDICADORES, CENTRO, REFERENCIA, self.user)
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "v2.xlsx"
            _workbook_v2(caminho)
            services.ingestar("ACOMPANHAMENTO_V2", caminho, CENTRO, REFERENCIA, self.user)
            services.ingestar("ACOMPANHAMENTO_V2", caminho, CENTRO, REFERENCIA, self.user)
        self.assertEqual(DespesaAcompanhamento.objects.count(), 2)  # substitui, não duplica
        self.assertEqual(ResumoFinanceiro.objects.count(), 13)        # demais fontes intactas
        self.assertEqual(ProjetoFinanceiro.objects.count(), 33)
        self.assertEqual(KpiAcompanhamento.objects.count(), 10)

    def test_v2_sem_metadados_importa_com_aviso_de_validacao_impossivel(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "sem_metadados.xlsx"
            _workbook_v2(caminho, com_metadados=False)
            imp = services.ingestar("ACOMPANHAMENTO_V2", caminho, CENTRO, REFERENCIA, self.user)
        self.assertEqual(imp.status, "SUCESSO")  # aviso, nunca erro
        self.assertEqual(imp.resumo, {"despesas": 2})
        self.assertIn(
            "Metadados de centro/período não encontrados no arquivo; "
            "validação contra o arquivo não foi possível.", imp.avisos)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(status="ERRO").count(), 0)
