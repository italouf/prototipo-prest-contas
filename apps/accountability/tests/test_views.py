from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from apps.core.permissions import adicionar_grupo, garantir_grupos


def _xlsx(bloco, nome="arquivo.xlsx"):
    return SimpleUploadedFile(nome, bloco, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class UploadPermissaoTestes(TestCase):
    def test_anonimo_redireciona_para_login(self):
        self.assertEqual(self.client.get(reverse("accountability:importar")).status_code, 302)

    def test_pontofocal_recebe_403(self):
        garantir_grupos()
        pf = adicionar_grupo(get_user_model().objects.create_user(username="pf", password="x"), "PontoFocal")
        self.client.force_login(pf)
        self.assertEqual(self.client.get(reverse("accountability:importar")).status_code, 403)


class UploadValidacaoTestes(TestCase):
    def setUp(self):
        garantir_grupos()
        self.master = adicionar_grupo(get_user_model().objects.create_user(username="m", password="x"), "Master")
        self.client.force_login(self.master)

    def test_extensao_csv_e_rejeitada(self):
        resposta = self.client.post(reverse("accountability:importar"), {
            "arquivo": SimpleUploadedFile("dados.csv", b"a", content_type="text/csv"),
            "centro": "QuIIN", "periodo_referencia": "2T/2024"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Envie um arquivo .xlsx")

    def test_centro_e_referencia_sao_obrigatorios(self):
        resposta = self.client.post(reverse("accountability:importar"), {
            "arquivo": _xlsx(b" "), "centro": "", "periodo_referencia": ""})
        self.assertContains(resposta, "obrigat")

    def test_sucesso_exibe_resumo_de_contagens(self):
        from pathlib import Path
        raiz = Path(__file__).resolve().parents[3]
        bloco = (raiz / "mockup" / "exemplos_arquivos" / "FINANCEIRO GERAL - REFAT.xlsx").read_bytes()
        resposta = self.client.post(reverse("accountability:importar"), {
            "arquivo": _xlsx(bloco, "FINANCEIRO GERAL - REFAT.xlsx"),
            "centro": "Centro de Competência Embrapii CIMATEC em Tecnologias Quânticas - Quiin",
            "periodo_referencia": "2T/2024"})
        self.assertContains(resposta, "33 projetos financeiros")
        self.assertContains(resposta, "2T/2024")

    def test_falha_de_parse_mostra_mensagem_amigavel(self):
        resposta = self.client.post(reverse("accountability:importar"), {
            "arquivo": _xlsx(b"PK\x03\x04lixo", "quebrado.xlsx"),
            "centro": "QuIIN", "periodo_referencia": "2T/2024"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(response=resposta, text="Não foi possível ler")


# --- Adições além dos testes verbatim ---------------------------------
# Os testes acima são o bloco do briefing, preservado literalmente. As
# adições abaixo (fixtures e testes extras) não alteram nenhum deles.

import io  # noqa: E402
from pathlib import Path  # noqa: E402

from openpyxl import Workbook  # noqa: E402

from apps.pillars.models import Pilar  # noqa: E402

from ..forms import UploadPrestacaoForm  # noqa: E402

# Pilares citados pelos arquivos reais: a ingestão rejeita pilar inexistente
# no banco (services._resolver_pilares — "Cadastre o pilar antes de aplicar"),
# então a fixture precisa existir para o upload de sucesso. É a mesma lista de
# apps/accountability/tests/test_ingestion.py.
_PILARES = (
    ("PDI", "PDI"), ("FORMACAO", "Formação FCRH"), ("STARTUPS", "ACS"),
    ("INFRA", "Infraestrutura"), ("AT", "AT"), ("OUTRASFONTES", "Outras Fontes"),
)


def _garantir_pilares():
    for codigo, nome in _PILARES:
        Pilar.objects.get_or_create(codigo=codigo, defaults={"nome": nome})


def _criar_pilares(cls):
    _garantir_pilares()


# Injetado sem editar o bloco verbatim: setUpTestData roda uma vez por classe
# antes de setUp e os dados vivem no mesmo nível de transação dos testes.
UploadValidacaoTestes.setUpTestData = classmethod(_criar_pilares)


def _xlsx_desconhecido(nome="desconhecido.xlsx"):
    """Workbook legível, mas sem assinatura de abas de nenhum template."""
    wb = Workbook()
    wb.active.title = "Qualquer coisa"
    buffer = io.BytesIO()
    wb.save(buffer)
    return _xlsx(buffer.getvalue(), nome)


class FormularioUploadTestes(TestCase):
    def test_periodo_mensal_aceita_meses_validos_e_remove_espacos(self):
        for referencia in ("01/2026", "09/2026", "12/2026", " 09/2026 "):
            with self.subTest(referencia=referencia):
                form = UploadPrestacaoForm(
                    data={"centro": "QuIIN", "periodo_referencia": referencia},
                    files={"arquivo": _xlsx(b"x")})
                self.assertTrue(form.is_valid(), form.errors)
                self.assertEqual(form.cleaned_data["periodo_referencia"], referencia.strip())

    def test_periodo_mensal_rejeita_mes_ano_e_formato_invalidos(self):
        for referencia in ("00/2026", "13/2026", "09/0000", "9/2026", "09/26", "09/20260",
                           "2026-09", "09/2026extra", "09/2026/01", "٠٩/٢٠٢٦"):
            with self.subTest(referencia=referencia):
                form = UploadPrestacaoForm(
                    data={"centro": "QuIIN", "periodo_referencia": referencia},
                    files={"arquivo": _xlsx(b"x")})
                self.assertFalse(form.is_valid())
                self.assertIn("periodo_referencia", form.errors)

    def test_periodo_normaliza_maiusculas_e_aceita_semestre(self):
        form = UploadPrestacaoForm(
            data={"centro": "QuIIN", "periodo_referencia": "2s/2024"},
            files={"arquivo": _xlsx(b"x")})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["periodo_referencia"], "2S/2024")

    def test_periodo_invalido_mostra_o_formato_esperado(self):
        form = UploadPrestacaoForm(
            data={"centro": "QuIIN", "periodo_referencia": "banana"},
            files={"arquivo": _xlsx(b"x")})
        self.assertFalse(form.is_valid())
        self.assertEqual(list(form.errors["periodo_referencia"]), [
            "Use MM/AAAA (ex.: 09/2026), trimestre (2T/2024) ou semestre (2S/2024)."])

    def test_extensao_e_validada_antes_de_qualquer_parse(self):
        form = UploadPrestacaoForm(
            data={"centro": "QuIIN", "periodo_referencia": "2T/2024"},
            files={"arquivo": _xlsx(b"a", "dados.ods")})
        self.assertFalse(form.is_valid())
        self.assertEqual(list(form.errors["arquivo"]), ["Envie um arquivo .xlsx."])

    def test_acima_de_25mb_e_rejeitado(self):
        form = UploadPrestacaoForm(
            data={"centro": "QuIIN", "periodo_referencia": "2T/2024"},
            files={"arquivo": _xlsx(b"a" * (25 * 1024 * 1024 + 1), "gigante.xlsx")})
        self.assertFalse(form.is_valid())
        self.assertIn("25 MB", form.errors["arquivo"][0])


class UploadExtrasTestes(TestCase):
    def setUp(self):
        garantir_grupos()
        _garantir_pilares()
        master = adicionar_grupo(
            get_user_model().objects.create_user(username="mx", password="x"), "Master")
        self.client.force_login(master)

    def test_get_renderiza_a_pagina_de_upload(self):
        resposta = self.client.get(reverse("accountability:importar"))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, 'data-testid="page-importar-prestacao"')
        self.assertContains(resposta, 'placeholder="09/2026"')
        self.assertContains(resposta, "Use MM/AAAA")

    def test_tipo_nao_identificado_mensagem_amigavel(self):
        resposta = self.client.post(reverse("accountability:importar"), {
            "arquivo": _xlsx_desconhecido(),
            "centro": "QuIIN", "periodo_referencia": "2T/2024"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "não foi possível identificar")

    def test_sucesso_grava_o_log_de_importacao(self):
        raiz = Path(__file__).resolve().parents[3]
        bloco = (raiz / "mockup" / "exemplos_arquivos" / "FINANCEIRO GERAL - REFAT.xlsx").read_bytes()
        self.client.post(reverse("accountability:importar"), {
            "arquivo": _xlsx(bloco, "FINANCEIRO GERAL - REFAT.xlsx"),
            "centro": "Centro de Competência Embrapii CIMATEC em Tecnologias Quânticas - Quiin",
            "periodo_referencia": "2T/2024"})
        from ..models import ImportacaoAcompanhamento
        importacao = ImportacaoAcompanhamento.objects.get(status="SUCESSO")
        self.assertEqual(importacao.resumo["projetos"], 33)
        self.assertEqual(importacao.tipo_fonte, "FINANCEIRO_GERAL")

    def test_importacao_mensal_reutiliza_acompanhamento_e_preserva_referencia_antiga(self):
        from ..models import Acompanhamento, ImportacaoAcompanhamento, KpiAcompanhamento, ResumoFinanceiro
        from .fixtures_financeiro import FINANCEIRO

        def enviar(caminho):
            resposta = self.client.post(reverse("accountability:importar"), {
                "arquivo": _xlsx(caminho.read_bytes(), caminho.name),
                "centro": "QuIIN", "periodo_referencia": "09/2026"})
            self.assertEqual(resposta.status_code, 200)
            self.assertContains(resposta, "09/2026")
            self.assertNotContains(resposta, 'data-testid="erro-periodo"')

        enviar(FINANCEIRO)
        mensal = Acompanhamento.objects.get(periodo_referencia="09/2026")
        self.assertTrue(ResumoFinanceiro.objects.filter(acompanhamento=mensal).exists())
        quantidade_resumos = ResumoFinanceiro.objects.filter(acompanhamento=mensal).count()
        antigo = Acompanhamento.objects.create(centro=mensal.centro, periodo_referencia="3T/2026")
        enviar(FINANCEIRO.with_name("Indicadores Gerais do Termo de Retificação do PE.xlsx"))
        self.assertEqual(KpiAcompanhamento.objects.filter(acompanhamento=mensal).count(), 10)
        enviar(FINANCEIRO)
        self.assertEqual(ResumoFinanceiro.objects.filter(acompanhamento=mensal).count(), quantidade_resumos)
        self.assertEqual(ImportacaoAcompanhamento.objects.filter(acompanhamento=mensal, status="SUCESSO").count(), 3)
        self.assertEqual(Acompanhamento.objects.filter(centro=mensal.centro).count(), 2)
        antigo.refresh_from_db()
        self.assertEqual(antigo.periodo_referencia, "3T/2026")

    def test_metadados_mensais_do_v2_precisam_corresponder_ao_upload(self):
        import tempfile
        from openpyxl import load_workbook
        from .test_ingestion import _workbook_v2, CENTRO
        from ..models import Acompanhamento, ImportacaoAcompanhamento

        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "acompanhamento.xlsx"
            _workbook_v2(caminho)
            wb = load_workbook(caminho)
            wb["0. Sumário"]["C3"] = "09/2026"
            wb.save(caminho)
            wb.close()
            bloco = caminho.read_bytes()
        for referencia, status in (("09/2026", "SUCESSO"), ("10/2026", "ERRO")):
            with self.subTest(referencia=referencia):
                resposta = self.client.post(reverse("accountability:importar"), {
                    "arquivo": _xlsx(bloco), "centro": CENTRO, "periodo_referencia": referencia})
                self.assertEqual(resposta.status_code, 200)
                self.assertTrue(ImportacaoAcompanhamento.objects.filter(status=status).exists())
                if status == "ERRO":
                    self.assertContains(resposta, "Metadado divergente")
        self.assertTrue(Acompanhamento.objects.filter(periodo_referencia="09/2026").exists())
        self.assertFalse(Acompanhamento.objects.filter(periodo_referencia="10/2026").exists())

    def test_formula_com_erro_importa_com_aviso_visivel_e_registrado(self):
        from io import BytesIO
        from .fixtures_financeiro import workbook_financeiro, tabela_financeira
        from ..models import ImportacaoAcompanhamento, ResumoFinanceiro

        wb = workbook_financeiro()
        ws, _tabela = tabela_financeira(wb, "tbl_ConsolidadoPrograma")
        ws["C2"] = "#REF!"
        bloco = BytesIO()
        wb.save(bloco)
        wb.close()
        resposta = self.client.post(reverse("accountability:importar"), {
            "arquivo": _xlsx(bloco.getvalue(), "nome-livre.xlsx"),
            "centro": "QuIIN", "periodo_referencia": "2T/2024"})
        self.assertContains(resposta, "33 projetos financeiros")
        self.assertContains(resposta, 'data-testid="avisos-importacao"')
        self.assertContains(resposta, "erro de fórmula")
        imp = ImportacaoAcompanhamento.objects.get(status="SUCESSO")
        self.assertIn("tbl_ConsolidadoPrograma", imp.avisos)
        self.assertIsNone(ResumoFinanceiro.objects.get(origem="TAB. 1", pilar__codigo="PDI").realizado)
