"""Contrato REFAT e resistência a mudanças de posição, ordem e tamanho."""
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableColumn

from apps.accountability.parsers.comum import normalizar, para_decimal
from apps.accountability.parsers.financeiro_geral import parse_financeiro_geral
from .fixtures_financeiro import (
    FINANCEIRO, preencher_tabela, tabela_financeira, workbook_financeiro,
)


def _resumo(payload, origem, pilar, ano=None):
    return next(r for r in payload["resumos"]
                if (r["origem"], r["pilar"], r["ano"]) == (origem, pilar, ano))


def _parse(wb):
    try:
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "nome-arbitrario.xlsx"
            wb.save(caminho)
            return parse_financeiro_geral(caminho)
    finally:
        wb.close()


def _celula(wb, tabela, cabecalho, registro=0):
    ws, objeto = tabela_financeira(wb, tabela)
    primeira, inicio, _ultima, _fim = range_boundaries(objeto.ref)
    indice = [c.name for c in objeto.tableColumns].index(cabecalho)
    return ws.cell(inicio + 1 + registro, primeira + indice)


class ParseFinanceiroGeralTestes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p = parse_financeiro_geral(FINANCEIRO)
        # A serialização do openpyxl arredonda floats do Excel além de 15
        # dígitos. Comparar alterações estruturais à mesma serialização base.
        cls.resalvo = _parse(workbook_financeiro())

    def test_valores_reais_e_contagens(self):
        self.assertEqual(self.p["erros"], [])
        self.assertEqual(len(self.p["resumos"]), 13)
        self.assertEqual(len(self.p["projetos"]), 33)
        for pilar, realizado in (("PDI", "17675109.56"), ("FORMACAO", "7749636.72"),
                                 ("STARTUPS", "2251069.32"), ("INFRA", "7299414.35")):
            self.assertEqual(_resumo(self.p, "TAB. 1", pilar)["realizado"], Decimal(realizado))
        self.assertEqual(_resumo(self.p, "TAB. 1", "PDI")["recurso_ou_meta"], Decimal("29000000"))
        self.assertEqual(sum(r["realizado"] for r in self.p["resumos"] if r["origem"] == "TAB. 1"), Decimal("34975229.95"))

    def test_reconciliacoes_e_captacao(self):
        for origem, pilar in (("TAB. 3", "PDI"), ("TAB. 5", "FORMACAO"), ("TAB. 8", "STARTUPS")):
            self.assertEqual(_resumo(self.p, origem, pilar)["realizado"], _resumo(self.p, "TAB. 1", pilar)["realizado"])
        at = _resumo(self.p, "TAB. 9", "AT")
        self.assertEqual(at["recurso_ou_meta"], Decimal("7750000"))
        self.assertEqual(at["captado"], Decimal("5616700"))
        self.assertEqual(at["captado_anos_1_2"], Decimal("346200"))
        self.assertEqual(at["realizado"], Decimal("220597"))
        self.assertEqual(_resumo(self.p, "TAB. 9", "OUTRASFONTES")["captado"], Decimal("10226768"))

    def test_projetos_sem_confundir_totais_e_campos_ausentes(self):
        projetos = self.p["projetos"]
        self.assertEqual([sum(p["origem"] == o for p in projetos) for o in ("TAB. 2", "TAB. 4", "TAB. 6")], [18, 8, 7])
        primeiro = projetos[0]
        self.assertEqual(primeiro["nome"], "Pós Processamento")
        self.assertEqual(primeiro["orcado"], Decimal("3799447.40"))
        self.assertEqual(primeiro["realizado"], Decimal("3795618.00"))
        self.assertEqual(primeiro["inicio"], date(2024, 1, 2))
        self.assertTrue(all(p["sequencia"] is None for p in projetos if p["pilar"] == "PDI"))
        self.assertTrue(all(p["inicio"] is None for p in projetos if p["pilar"] == "STARTUPS"))
        self.assertFalse(any(p["nome"].startswith("TOTAL") for p in projetos))

    def test_anuais_vazios_preservam_resumos_e_emitem_aviso(self):
        self.assertFalse(any(r["ano"] is not None for r in self.p["resumos"]))
        self.assertEqual(sum(r["origem"] == "TAB. 1.1" for r in self.p["resumos"]), 4)
        self.assertTrue(any("tbl_VisaoGeral" in a for a in self.p["avisos"]))
        self.assertFalse(any("TAB. 7" in a for a in self.p["avisos"]))

    def test_normalizacao_numerica(self):
        self.assertEqual(para_decimal("1.234,56"), Decimal("1234.56"))
        self.assertIsNone(para_decimal("#REF!"))
        self.assertEqual(normalizar("  AFCCT / PD&I "), "afcct / pd&i")

    def test_renomear_e_reordenar_abas_nao_muda_payload(self):
        wb = workbook_financeiro()
        for indice, ws in enumerate(wb):
            ws.title = f"Dados {indice}"
        wb.move_sheet(wb.worksheets[-1], offset=-4)
        wb.create_sheet("Anotações")["A1"] = "Conteúdo fora das tabelas"
        self.assertEqual(_parse(wb), self.resalvo)

    def test_mover_todas_as_tabelas_para_outra_aba(self):
        wb = workbook_financeiro()
        destino = wb.create_sheet("Objetos reorganizados")
        linha_inicio = 10
        for ws in list(wb.worksheets[:-1]):
            for tabela in list(ws.tables.values()):
                dados = [[c.value for c in linha] for linha in ws[tabela.ref]]
                del ws.tables[tabela.displayName]
                for i, linha in enumerate(dados, start=linha_inicio):
                    for j, valor in enumerate(linha, start=5):
                        destino.cell(i, j, valor)
                tabela.ref = f"E{linha_inicio}:{get_column_letter(4 + len(dados[0]))}{linha_inicio + len(dados) - 1}"
                destino.add_table(tabela)
                linha_inicio += len(dados) + 10
        self.assertEqual(_parse(wb), self.resalvo)

    def test_inverter_colunas_de_todas_as_tabelas(self):
        wb = workbook_financeiro()
        for ws in wb:
            for tabela in ws.tables.values():
                for linha in ws[tabela.ref]:
                    dados = [(c.value, c.number_format) for c in linha][::-1]
                    for celula, (valor, formato) in zip(linha, dados):
                        celula.value = valor
                        celula.number_format = formato
                tabela.tableColumns.reverse()
        self.assertEqual(_parse(wb), self.resalvo)

    def test_coluna_auxiliar_e_tabela_extra_sao_ignoradas(self):
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        primeira, inicio, ultima, fim = range_boundaries(tabela.ref)
        for linha in range(inicio, fim + 1):
            ws.cell(linha, ultima + 1, "OBSERVAÇÃO" if linha == inicio else "auxiliar")
        tabela.tableColumns.append(TableColumn(id=99, name="OBSERVAÇÃO"))
        tabela.ref = f"{get_column_letter(primeira)}{inicio}:{get_column_letter(ultima + 1)}{fim}"
        extra = wb.create_sheet("Notas")
        extra.append(["Chave", "Valor"])
        extra.append(["Nota", "Texto"])
        extra.add_table(Table(displayName="tbl_Notas", ref="A1:B2"))
        self.assertEqual(_parse(wb), self.resalvo)

    def test_incluir_projeto_apos_total_e_expandir_objeto(self):
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        primeira, inicio, ultima, fim = range_boundaries(tabela.ref)
        ws.cell(fim + 1, primeira, "Projeto Novo")
        ws.cell(fim + 1, primeira + 1, "Em andamento")
        ws.cell(fim + 1, primeira + 4, 123.45)
        tabela.ref = f"{get_column_letter(primeira)}{inicio}:{get_column_letter(ultima)}{fim + 1}"
        resultado = _parse(wb)
        self.assertEqual(resultado["erros"], [])
        self.assertEqual(len(resultado["projetos"]), 34)
        self.assertEqual(resultado["projetos"][18]["nome"], "Projeto Novo")

    def test_conteudo_fora_do_objeto_nao_vira_projeto(self):
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        primeira, _inicio, _ultima, fim = range_boundaries(tabela.ref)
        ws.cell(fim + 1, primeira, "Projeto fora da tabela")
        self.assertEqual(_parse(wb), self.resalvo)

    def test_totais_nativos_e_nomes_com_total(self):
        wb = workbook_financeiro()
        ws, tabela = preencher_tabela(wb, "tbl_ProjetosPDI", [
            {"PROJETO": "Total de aplicações quânticas"},
            {"PROJETO": "TOTAL GERAL"},
            {"PROJETO": "Projeto após total"},
            {"PROJETO": "#Projeto identificado", "STATUS": "#Em andamento"},
            {"PROJETO": "Rodapé nativo"},
        ])
        tabela.totalsRowCount = 1
        tabela.totalsRowShown = True
        resultado = _parse(wb)
        self.assertEqual(resultado["erros"], [])
        self.assertEqual([p["nome"] for p in resultado["projetos"] if p["pilar"] == "PDI"],
                         ["Total de aplicações quânticas", "Projeto após total", "#Projeto identificado"])

    def test_campos_anuais_incluem_2027_sem_transformar_total_em_ano(self):
        wb = workbook_financeiro()
        preencher_tabela(wb, "tbl_VisaoGeral", [{
            "AÇÃO": "AFCCT / PD&I", "RECURSO PPI TOTAL": 29000000,
            "REALIZADO 2024": 111.11, "REALIZADO 2025": 222.22,
            "REALIZADO 2026 (YTD)": 333.33, "PROJETADO 2026": 444.44,
            "PROJETADO 2027": 0, "PROJETADO TOTAL": 777.77,
        }])
        resultado = _parse(wb)
        self.assertEqual(resultado["erros"], [])
        self.assertEqual(_resumo(resultado, "TAB. 1.1", "PDI", 2024)["realizado"], Decimal("111.11"))
        self.assertEqual(_resumo(resultado, "TAB. 1.1", "PDI", 2026)["projetado"], Decimal("444.44"))
        self.assertEqual(_resumo(resultado, "TAB. 1.1", "PDI", 2027)["projetado"], Decimal("0.00"))
        self.assertIsNone(_resumo(resultado, "TAB. 1.1", "PDI", 2027)["realizado"])
        self.assertEqual(_resumo(resultado, "TAB. 1.1", "PDI")["projetado"], Decimal("777.77"))

    def test_atualizar_mes_e_normalizar_cabecalhos(self):
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        primeira, inicio, _ultima, _fim = range_boundaries(tabela.ref)
        for indice, coluna in enumerate(tabela.tableColumns):
            nome = coluna.name.replace("05/2026", "06/2026").replace("06 a 12/2026", "07 a 12/2026")
            nome = "  " + nome.lower().replace(" ", "  ") + "  "
            coluna.name = nome
            ws.cell(inicio, primeira + indice).value = nome
        ws, tabela = tabela_financeira(wb, "tbl_VisaoPDI")
        tabela.tableColumns[1].name = "RECURSO PPI_x000a_AFCCT / PDI"
        self.assertEqual(_parse(wb), self.resalvo)

    def test_erros_de_formula_e_cache_ausente_viram_avisos(self):
        for valor, trecho in (("#REF!", "erro de fórmula"), ("=SUM(1,2)", "sem resultado salvo")):
            with self.subTest(valor=valor):
                wb = workbook_financeiro()
                _celula(wb, "tbl_ConsolidadoPrograma", "REALIZADO TOTAL").value = valor
                resultado = _parse(wb)
                self.assertEqual(resultado["erros"], [])
                self.assertIsNone(_resumo(resultado, "TAB. 1", "PDI")["realizado"])
                self.assertTrue(any(trecho in a and "tbl_ConsolidadoPrograma" in a and "realizado" in a for a in resultado["avisos"]))

    def test_validacao_estrutural_e_de_identidade(self):
        for problema in ("tabela ausente", "coluna ausente", "metadados", "duplicado", "pilar"):
            with self.subTest(problema=problema):
                wb = workbook_financeiro()
                ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
                if problema == "tabela ausente":
                    del ws.tables[tabela.displayName]
                elif problema == "coluna ausente":
                    tabela.tableColumns[0].name = "OUTRO CAMPO"
                    ws["A1"] = "OUTRO CAMPO"
                elif problema == "metadados":
                    ws["A1"] = "ALTERAÇÃO SEM METADADOS"
                elif problema == "duplicado":
                    _celula(wb, "tbl_ProjetosPDI", "PROJETO", 1).value = "Pós Processamento"
                else:
                    _celula(wb, "tbl_ConsolidadoPrograma", "AÇÃO").value = "PILAR INVÁLIDO"
                self.assertTrue(_parse(wb)["erros"])

    def test_colunas_ambiguas_sao_rejeitadas(self):
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_ProjetosPDI")
        tabela.tableColumns[10].name = "REALIZADO TOTAL"
        ws["K1"] = "REALIZADO TOTAL"
        self.assertTrue(any("ambíguo" in e for e in _parse(wb)["erros"]))

    def test_valores_incompativeis_nao_sao_zeros_silenciosos(self):
        for cabecalho, valor in (("ORÇADO SÍNTESE", "texto"), ("INICIO", "data inválida"),
                                 ("ORÇADO SÍNTESE", "NaN"), ("PROJETO", "x" * 256)):
            with self.subTest(cabecalho=cabecalho, valor=valor):
                wb = workbook_financeiro()
                _celula(wb, "tbl_ProjetosPDI", cabecalho).value = valor
                self.assertTrue(_parse(wb)["erros"])

    def test_formato_antigo_nao_e_aceito(self):
        antigo = FINANCEIRO.with_name("FINANCEIRO GERAL.xlsx")
        self.assertTrue(parse_financeiro_geral(antigo)["erros"])
