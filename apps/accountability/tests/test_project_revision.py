from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from apps.core.permissions import adicionar_grupo, garantir_grupos
from apps.pillars.models import Pilar
from apps.planning.views import contexto_pilar, _html_standalone
from apps.core.pillar_overview import grafico_carteira

from ..models import Acompanhamento, ProjetoFinanceiro
from .. import services
from ..parsers.financeiro_geral import parse_financeiro_geral
from ..project_status import STATUS_PROJETOS, etapa_do_status
from ..refat_workbook import revisar_planilha
from .fixtures_financeiro import FINANCEIRO


class RevisaoPlanilhaTestes(SimpleTestCase):
    def test_revisao_preserva_valores_datas_formulas_e_aceita_novos_campos(self):
        original = FINANCEIRO.read_bytes()
        antes = parse_financeiro_geral(BytesIO(original))
        with TemporaryDirectory() as directory:
            arquivo = Path(directory) / "refat.xlsx"
            arquivo.write_bytes(original)
            resultado = revisar_planilha(arquivo)
            self.assertTrue(resultado["alterado"])
            depois = parse_financeiro_geral(arquivo)
            self.assertEqual(depois["erros"], [])
            self.assertEqual(antes["resumos"], depois["resumos"])
            self.assertEqual(len(depois["projetos"]), 33)
            for a, b in zip(antes["projetos"], depois["projetos"]):
                for campo in ("nome", "status", "inicio", "fim", "orcado", "realizado", "projetado_2026", "projetado_2027", "realizado_mais_projetado", "diferenca"):
                    self.assertEqual(a[campo], b[campo], (a["nome"], campo))
            wb = load_workbook(BytesIO(arquivo.read_bytes()), data_only=False)
            for nome in ("tbl_ProjetosPDI", "tbl_ProjetosFCRH", "tbl_ProjetosACS"):
                sheet = next(s for s in wb if nome in s.tables)
                headers = [c.name for c in sheet.tables[nome].tableColumns]
                for header in ("ASSINATURA DA SÍNTESE", "FINAL PREVISTO NA SÍNTESE", "DATA DO KICK-OFF", "FINAL PREVISTO REAL", "FAROL", "ETAPA DO PROJETO"):
                    self.assertIn(header, headers)
                self.assertEqual(sheet.data_validations.dataValidation[0].formula1, "StatusProjetos")
                self.assertTrue(any(c.data_type == "f" for row in sheet for c in row))
            opcoes = [r[0].value for r in wb["_Opcoes_Projetos"].iter_rows(min_row=2)]
            for status in STATUS_PROJETOS:
                self.assertIn(status, opcoes)
            wb.close()
            self.assertFalse(revisar_planilha(arquivo)["alterado"])

    def test_novos_status_preservam_diferenca_entre_encerrados_e_aprovacoes(self):
        self.assertEqual(etapa_do_status("Encerrado apenas no CIMATEC"), "Encerrados")
        self.assertEqual(etapa_do_status("Encerrado formalmente EMBRAPII"), "Encerrados")
        for status in ("Em aprovação EMBRAPII", "Em aprovação interna (pré-projeto)", "Em elaboração"):
            self.assertEqual(etapa_do_status(status), "Futuros projetos")


class ProjetoFinanceiroRevisadoTestes(TestCase):
    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        cls.user = adicionar_grupo(get_user_model().objects.create_user(username="projetos-revisados"), "Master")
        cls.pilares = {c: Pilar.objects.get_or_create(codigo=c, defaults={"nome": c})[0] for c in ("PDI", "FORMACAO", "STARTUPS", "INFRA", "AT", "OUTRASFONTES")}
        services.ingestar("FINANCEIRO_GERAL", FINANCEIRO, "Projetos revisados", "3T/2026", cls.user)
        cls.acomp = Acompanhamento.objects.get(centro__nome="Projetos revisados")

    def painel(self, codigo):
        return contexto_pilar(self.pilares[codigo], "todos", "fin", self.user, self.acomp)["painel_pilar"]

    def test_cards_na_ordem_solicitada_e_modelo_replicado(self):
        for codigo, titulo in (("PDI", "Recurso PPI AFCCT total"), ("FORMACAO", "Recurso PPI FCRH total"), ("STARTUPS", "Recurso PPI ACS total")):
            with self.subTest(codigo=codigo):
                pp = self.painel(codigo)
                self.assertEqual([c["titulo"] for c in pp["cards_overview"]], [titulo, "Recurso PPI realizado", "Diferença rel. PPI"])
                self.assertTrue(pp["cards_overview"][1]["hero"])
                self.assertEqual(sum(g["quantidade"] for g in pp["grupos_projetos"]), pp["projetos"])
                self.assertNotIn("Em preparação", [g["titulo"] for g in pp["grupos_projetos"]])
                self.assertTrue(pp["grafico_financeiro"]["linhas"])

    def test_datas_status_e_farol_passam_pela_importacao_e_exportacao(self):
        with TemporaryDirectory() as directory:
            arquivo = Path(directory)/"revisado.xlsx"
            arquivo.write_bytes(FINANCEIRO.read_bytes())
            revisar_planilha(arquivo)
            wb = load_workbook(BytesIO(arquivo.read_bytes()), data_only=True)
            ws = wb["db_afcct_pdi"]
            headers = [c.name for c in ws.tables["tbl_ProjetosPDI"].tableColumns]
            ws.cell(2, headers.index("STATUS")+1, "Em aprovação EMBRAPII")
            ws.cell(2, headers.index("ETAPA DO PROJETO")+1, "Futuros projetos")
            ws.cell(2, headers.index("DATA DO KICK-OFF")+1, date(2026, 9, 1))
            ws.cell(2, headers.index("FINAL PREVISTO REAL")+1, date(2027, 6, 30))
            wb.save(arquivo)
            wb.close()
            services.ingestar("FINANCEIRO_GERAL", arquivo, "Projetos revisados", "3T/2026", self.user)
        projeto = ProjetoFinanceiro.objects.get(acompanhamento=self.acomp, nome="Pós Processamento")
        self.assertEqual(projeto.data_kickoff, date(2026, 9, 1))
        self.assertEqual(projeto.fim_previsto_real, date(2027, 6, 30))
        self.assertEqual(projeto.inicio, date(2024, 1, 2))
        self.assertEqual(projeto.etapa, "Futuros projetos")
        self.assertEqual(projeto.farol, "Meta atingida")
        self.client.force_login(self.user)
        response = self.client.get(f"/pilar/{self.pilares['PDI'].pk}/?base=fin")
        self.assertEqual(response.status_code, 200)
        for texto in ("Assinatura da síntese", "Data do kick-off", "Final previsto real", "Em aprovação EMBRAPII", "Futuros projetos", 'data-testid="grafico-financeiro-pilar"'):
            self.assertIn(texto, response.content.decode())
        self.assertNotIn('class="overview-table pillar-financial-table"', response.content.decode())
        html = _html_standalone("todos", "fin", self.user, pilar=self.pilares["PDI"], acompanhamento=self.acomp)
        self.assertIn("30/06/2027", html)
        self.assertIn('data-testid="grafico-financeiro-pilar"', html)

    def test_grafico_nao_inventa_projecoes_e_respeita_escala(self):
        pp = self.painel("STARTUPS")
        for linha in pp["grafico_financeiro"]["linhas"]:
            self.assertLessEqual(linha["largura_orcado"], 100)
            self.assertLessEqual(sum(s["largura"] for s in linha["segmentos"]), Decimal("100.000000001"))
            self.assertIsNone(linha["segmentos"][1]["valor"])
            self.assertIsNone(linha["segmentos"][2]["valor"])
        self.assertTrue(grafico_carteira([])["sem_dados"])
