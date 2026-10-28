"""Metas financeiras da Visão Geral: cálculos, edição e importação compatível."""
from decimal import Decimal as D
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import SimpleTestCase, TestCase

from apps.audit.models import AuditLog
from apps.core.permissions import adicionar_grupo, garantir_grupos
from apps.pillars.models import Pilar
from apps.planning.views import contexto_painel, _html_standalone

from .. import overrides
from ..models import Acompanhamento, CentroCompetencia, ResumoFinanceiro
from ..parsers.financeiro_geral import parse_financeiro_geral
from .fixtures_financeiro import workbook_financeiro, tabela_financeira


class OverviewMetasTestes(TestCase):
    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        cls.usuario = adicionar_grupo(
            get_user_model().objects.create_user(username="overview-metas", password="test-123"), "Master")
        cls.pilares = {}
        for codigo in ("PDI", "FORMACAO", "STARTUPS", "INFRA", "AT", "OUTRASFONTES"):
            cls.pilares[codigo], _ = Pilar.objects.get_or_create(codigo=codigo, defaults={"nome": codigo})
        centro = CentroCompetencia.objects.create(codigo="overview-metas", nome="Overview Metas")
        cls.acomp = Acompanhamento.objects.create(centro=centro, periodo_referencia="3T/2026")
        for codigo, recurso, realizado in (
                ("PDI", "29000000", "17675109.56"),
                ("FORMACAO", "14000000", "7749636.72"),
                ("STARTUPS", "7000000", "2251069.32"),
                ("INFRA", "10000000", "7299414.35")):
            ResumoFinanceiro.objects.create(acompanhamento=cls.acomp, origem="TAB. 1",
                pilar=cls.pilares[codigo], recurso_ou_meta=D(recurso), realizado=D(realizado))
        for codigo, meta, captado, realizado in (
                ("AT", "7750000", "5616700", "220597"),
                ("OUTRASFONTES", "6300000", "10226768", "242112")):
            ResumoFinanceiro.objects.create(acompanhamento=cls.acomp, origem="TAB. 9",
                pilar=cls.pilares[codigo], recurso_ou_meta=D(meta), captado=D(captado),
                realizado=D(realizado))

    def painel(self, selecionados=None, periodo=None):
        return contexto_painel("todos", "fin", self.usuario, self.acomp, selecionados, periodo)["painel"]

    def test_seis_cards_seguem_a_ordem_e_usam_metas_sem_alterar_contrato_compartilhado(self):
        painel = self.painel()
        cards = painel["cards_resumo"]
        self.assertEqual([c["chave"] for c in cards],
                         ["ppi", "at", "at-captado", "outras", "outras-captado", "total"])
        self.assertEqual(cards[1]["previsto"], D("7750000"))
        self.assertEqual(cards[2]["executado"], D("5616700"))
        self.assertEqual(cards[3]["previsto"], D("6300000"))
        self.assertEqual(cards[-1]["previsto"], D("74050000"))
        self.assertEqual(cards[-1]["executado"], D("35437938.95"))
        self.assertEqual(len(painel["cards"]), 4)

    def test_quadros_independentes_consolidam_apenas_as_fontes(self):
        captacao, execucao = self.painel()["tabelas_fontes"]
        self.assertEqual(captacao["total"]["recurso"], D("14050000"))
        self.assertEqual(captacao["total"]["realizado"], D("15843468"))
        self.assertEqual(captacao["total"]["diferenca"], D("-1793468"))
        self.assertEqual(execucao["total"]["realizado"], D("462709"))
        self.assertEqual(execucao["total"]["diferenca"], D("13587291"))
        self.assertEqual(execucao["total"]["pct"], D("462709") / D("14050000") * 100)
        outras = captacao["linhas"][1]
        self.assertGreater(outras["pct"], 100)
        self.assertEqual(outras["barra"], 100)
        self.assertEqual(outras["faixa"], "ok")

    def test_selecao_de_fontes_nao_vaza_valores_excluidos(self):
        painel = self.painel(["AT"])
        for tabela in painel["tabelas_fontes"]:
            self.assertEqual(len(tabela["linhas"]), 1)
            self.assertEqual(tabela["total"]["recurso"], D("7750000"))
        self.assertEqual(painel["cards_resumo"][-1]["previsto"], D("7750000"))
        self.assertEqual(painel["cards_resumo"][-1]["executado"], D("220597"))
        self.assertIsNone(painel["cards_resumo"][3]["executado"])
        for selecionados in ([], ["PDI"]):
            for tabela in self.painel(selecionados)["tabelas_fontes"]:
                self.assertEqual(tabela["linhas"], [])
                self.assertIsNone(tabela["total"]["recurso"])
                self.assertIsNone(tabela["total"]["pct"])

    def test_ausencia_zero_e_meta_zero_sao_distintos(self):
        at = ResumoFinanceiro.objects.get(acompanhamento=self.acomp, pilar=self.pilares["AT"])
        for meta, captado, esperado in ((None, None, None), (D("100"), D("0"), D("0")),
                                        (D("0"), D("50"), None)):
            with self.subTest(meta=meta, captado=captado):
                at.recurso_ou_meta, at.captado = meta, captado
                at.save()
                linha = self.painel(["AT"])["tabelas_fontes"][0]["linhas"][0]
                self.assertEqual(linha["pct"], esperado)
                self.assertEqual(linha["faixa"], "critica" if esperado == 0 else "neutra")

    def test_farol_considera_precisao_real_e_nao_arredonda_a_faixa(self):
        at = ResumoFinanceiro.objects.get(acompanhamento=self.acomp, pilar=self.pilares["AT"])
        for captado, faixa in (("49.99", "critica"), ("50", "parcial"),
                                ("89.99", "parcial"), ("90", "ok"), ("150", "ok")):
            with self.subTest(captado=captado):
                at.recurso_ou_meta, at.captado = D("100"), D(captado)
                at.save()
                self.assertEqual(self.painel(["AT"])["tabelas_fontes"][0]["linhas"][0]["faixa"], faixa)

    def test_intervalo_sem_dados_limpa_fontes_cards_e_metas_anuais(self):
        from apps.core.overview_period import resolver_periodo
        ResumoFinanceiro.objects.create(acompanhamento=self.acomp, origem="TAB. 1.1",
            pilar=self.pilares["PDI"], ano=2024, recurso_ou_meta=D("100"), realizado=D("50"))
        for params in ({"mes_inicial": "02/2025", "mes_final": "03/2025"},
                       {"mes_inicial": "13/2025", "mes_final": "03/2025"}):
            painel = self.painel(periodo=resolver_periodo(params))
            for card in painel["cards_resumo"]:
                self.assertIsNone(card["executado"])
                self.assertIsNone(card["previsto"])
            self.assertIsNone(painel["tabela_tab11"]["linhas"][0]["meta_2024"])
            self.assertTrue(painel["tabela_tab11"]["sem_dados_anuais"])
            for tabela in painel["tabelas_fontes"]:
                self.assertIsNone(tabela["total"]["recurso"])
                self.assertIsNone(tabela["total"]["realizado"])

    def test_farol_anual_usa_realizado_sem_depender_de_metas_ou_projecoes(self):
        for ano, realizado in ((2024, "1000000"), (2025, "2000000"), (2026, "3000000")):
            ResumoFinanceiro.objects.create(acompanhamento=self.acomp, origem="TAB. 1.1",
                pilar=self.pilares["PDI"], ano=ano, recurso_ou_meta=D("100000"),
                realizado=D(realizado), projetado=D("20000000") if ano == 2026 else None)
        linha = self.painel()["tabela_tab11"]["linhas"][0]
        self.assertEqual(linha["pct"], D("6000000") / D("29000000") * 100)
        self.assertEqual(linha["faixa"], "critica")
        self.assertEqual(linha["combinado"], D("26000000"))
        self.assertEqual(linha["diferenca"], D("3000000"))
        self.assertEqual(linha["meta_2024"], D("100000"))

    def test_meta_anual_e_captacao_editam_referencias_corretas_com_auditoria(self):
        ResumoFinanceiro.objects.create(acompanhamento=self.acomp, origem="TAB. 1.1",
            pilar=self.pilares["PDI"], ano=2024, recurso_ou_meta=D("123"))
        overrides.aplicar_overrides(self.acomp, "financeiro", [
            (2024, "PDI", "meta_anual", D("456")),
            (0, "PDI", "previsto", D("30000000")),
            (0, "AT", "meta_captacao", D("8000000")),
        ], self.usuario)
        painel = self.painel()
        self.assertEqual(painel["tabela_tab11"]["linhas"][0]["meta_2024"], D("456"))
        self.assertEqual(painel["tabela_tab11"]["linhas"][0]["recurso"], D("30000000"))
        self.assertEqual(painel["cards_resumo"][1]["previsto"], D("8000000"))
        self.assertTrue(AuditLog.objects.filter(campo="meta_anual", valor_novo="456").exists())
        overrides.aplicar_overrides(self.acomp, "financeiro", [(2024, "PDI", "meta_anual", None)], self.usuario)
        self.assertEqual(self.painel()["tabela_tab11"]["linhas"][0]["meta_2024"], D("123"))

    def test_meta_anual_valida_escopo_anos_precisao_e_permissoes(self):
        for base, ano, codigo, valor in (
                ("financeiro", 0, "PDI", D("1")), ("financeiro", 2027, "PDI", D("1")),
                ("financeiro", 2024, "AT", D("1")), ("fisico", 2024, "PE-01", D("1")),
                ("financeiro", 2024, "PDI", D("-1")), ("financeiro", 2024, "PDI", D("1.001"))):
            with self.subTest(base=base, ano=ano, codigo=codigo, valor=valor), self.assertRaises(ValidationError):
                overrides.aplicar_overrides(self.acomp, base, [(ano, codigo, "meta_anual", valor)], self.usuario)
        leitor = adicionar_grupo(get_user_model().objects.create_user(username="metas-leitor"), "Lideranca")
        with self.assertRaises(PermissionDenied):
            overrides.aplicar_overrides(self.acomp, "fin", [(2024, "PDI", "meta_anual", D("1"))], leitor)

    def test_pagina_fragmento_e_exportacao_incluem_nova_apresentacao(self):
        self.client.force_login(self.usuario)
        url = f"/?ano=todos&base=fin&acompanhamento={self.acomp.pk}"
        for headers in ({}, {"HTTP_HX_REQUEST": "true"}):
            resposta = self.client.get(url, **headers)
            self.assertEqual(resposta.status_code, 200)
            html = resposta.content.decode()
            self.assertEqual(html.count(' id="dashboard-controls"'), 1)
            self.assertEqual(html.count('data-testid="base-toggle"'), 1)
            self.assertIn("Visão Consolidada do PPI", html)
            self.assertIn("Meta ano 1 - 2024", html)
            self.assertIn('data-testid="fontes-captado"', html)
            self.assertIn('name="v__2024__financeiro__PDI__meta_anual"', html)
            self.assertNotIn("recurso EMBRAPII", html)
        html = _html_standalone("todos", "fin", self.usuario, acompanhamento=self.acomp)
        self.assertIn("Meta ano 3 - 2026", html)
        self.assertIn('data-testid="fontes-realizado"', html)
        self.assertNotIn(' id="dashboard-controls"', html)


class ImportacaoMetasAnuaisTestes(SimpleTestCase):
    def test_planilha_com_metas_e_novos_cabecalhos_preserva_realizados_e_projecoes(self):
        from openpyxl.worksheet.table import TableColumn
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        wb = workbook_financeiro()
        ws, tabela = tabela_financeira(wb, "tbl_VisaoGeral")
        primeira, inicio, ultima, fim = range_boundaries(tabela.ref)
        renomes = {
            "REALIZADO 2024": "REALIZADO ANO 1 - 2024",
            "REALIZADO 2025": "REALIZADO ANO 2 - 2025",
            "REALIZADO 2026 (YTD)": "REALIZADO ANO 3 (YTD) - 2026",
        }
        for indice, coluna in enumerate(tabela.tableColumns, start=primeira):
            if coluna.name in renomes:
                coluna.name = renomes[coluna.name]
                ws.cell(inicio, indice, coluna.name)
        for indice, ano in enumerate((2024, 2025, 2026), start=1):
            numero = ultima + indice
            cabecalho = f"META ANO {indice} - {ano}"
            ws.cell(inicio, numero, cabecalho)
            ws.cell(inicio + 1, numero, 0 if ano == 2025 else 1000000 * indice)
            tabela.tableColumns.append(TableColumn(id=numero-primeira+1, name=cabecalho))
        tabela.ref = f"{get_column_letter(primeira)}{inicio}:{get_column_letter(ultima+3)}{fim}"
        with TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "metas.xlsx"
            wb.save(arquivo)
            payload = parse_financeiro_geral(arquivo)
        wb.close()
        self.assertEqual(payload["erros"], [])
        anuais = {r["ano"]: r for r in payload["resumos"] if r["origem"] == "TAB. 1.1" and r["pilar"] == "PDI" and r["ano"]}
        self.assertEqual(anuais[2024]["recurso_ou_meta"], D("1000000"))
        self.assertEqual(anuais[2025]["recurso_ou_meta"], D("0"))
        self.assertEqual(anuais[2026]["recurso_ou_meta"], D("3000000"))
        self.assertIsNone(anuais[2024]["realizado"])
        self.assertIsNone(anuais[2024]["projetado"])
