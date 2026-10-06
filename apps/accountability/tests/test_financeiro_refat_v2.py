"""Importação do REFAT v2, códigos EMBRAPII e acesso aos dashboards."""
import tempfile
from decimal import Decimal
from pathlib import Path

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from apps.core.permissions import adicionar_grupo, garantir_grupos
from apps.pillars.models import Pilar

from .. import services
from ..models import Acompanhamento, ProjetoFinanceiro
from .fixtures_financeiro import FINANCEIRO, tabela_financeira


FINANCEIRO_V2 = FINANCEIRO.with_name("FINANCEIRO GERAL - REFAT - v2.xlsx")


class FinanceiroRefatV2Testes(TestCase):
    @classmethod
    def setUpTestData(cls):
        garantir_grupos()
        cls.usuario = adicionar_grupo(
            get_user_model().objects.create_user(username="refat-v2"), "Master")
        cls.pilares = {
            codigo: Pilar.objects.create(codigo=codigo, nome=codigo)
            for codigo in ("PDI", "FORMACAO", "STARTUPS", "INFRA", "AT", "OUTRASFONTES")
        }

    def importar(self, caminho=FINANCEIRO_V2):
        return services.ingestar(
            "FINANCEIRO_GERAL", caminho, "Centro QuIIN", "2T/2024", self.usuario)

    def test_arquivo_corrigido_importa_sem_tratar_totais_como_projetos(self):
        importacao = self.importar()
        self.assertEqual(importacao.status, "SUCESSO")
        self.assertEqual(importacao.resumo, {"resumos": 13, "projetos": 33})
        for codigo, quantidade in (("PDI", 18), ("FORMACAO", 8), ("STARTUPS", 7)):
            self.assertEqual(ProjetoFinanceiro.objects.filter(pilar__codigo=codigo).count(), quantidade)
        fcrh = ProjetoFinanceiro.objects.filter(pilar__codigo="FORMACAO").order_by("pk").first()
        self.assertEqual(fcrh.codigo_projeto_embrapii, "1")
        self.assertEqual(fcrh.realizado, Decimal("39402.73"))

        self.client.force_login(self.usuario)
        for codigo, nome in (("PDI", "PDI"), ("FORMACAO", "FCRH"), ("STARTUPS", "ACS")):
            with self.subTest(codigo=codigo):
                resposta = self.client.get(reverse("core:pilar", args=[self.pilares[codigo].pk]))
                self.assertContains(resposta, "CÓDIGO DO PROJETO EMBRAPII")
                self.assertContains(resposta, f'data-testid="tbl_Projetos{nome}"')
                self.assertContains(resposta, f'data-testid="tbl_Visao{nome}"')

    def test_codigos_alfanumericos_sao_preservados_na_gravacao_e_exportacao(self):
        codigos = {"PDI": "001-PDI/2026", "FORMACAO": "002-FCRH/2026", "STARTUPS": "003-ACS/2026"}
        wb = load_workbook(FINANCEIRO_V2, data_only=True)
        try:
            for codigo, nome in (("PDI", "PDI"), ("FORMACAO", "FCRH"), ("STARTUPS", "ACS")):
                ws, tabela = tabela_financeira(wb, f"tbl_Projetos{nome}")
                primeira, inicio, _ultima, _fim = range_boundaries(tabela.ref)
                indice = [coluna.name for coluna in tabela.tableColumns].index("CÓDIGO DO PROJETO EMBRAPII")
                ws.cell(inicio + 1, primeira + indice).value = codigos[codigo]
            with tempfile.TemporaryDirectory() as pasta:
                caminho = Path(pasta) / "codigos.xlsx"
                wb.save(caminho)
                self.importar(caminho)
                self.importar(caminho)
        finally:
            wb.close()
        self.assertEqual(ProjetoFinanceiro.objects.count(), 33)
        self.client.force_login(self.usuario)
        for codigo, identificador in codigos.items():
            projeto = ProjetoFinanceiro.objects.filter(pilar__codigo=codigo).order_by("pk").first()
            self.assertEqual(projeto.codigo_projeto_embrapii, identificador)
            resposta = self.client.get(reverse("core:pilar", args=[self.pilares[codigo].pk]))
            self.assertContains(resposta, identificador)
            for destino in ("planning:csv", "planning:export_html"):
                resposta = self.client.get(reverse(destino), {"pilar": self.pilares[codigo].pk, "base": "fin"})
                self.assertContains(resposta, identificador)

    def test_coluna_de_realizado_nao_pode_ser_confundida_com_codigo(self):
        self.importar()
        antes = list(ProjetoFinanceiro.objects.order_by("pk").values_list("nome", "realizado"))
        wb = load_workbook(FINANCEIRO_V2, data_only=True)
        try:
            ws, tabela = tabela_financeira(wb, "tbl_ProjetosFCRH")
            primeira, inicio, _ultima, _fim = range_boundaries(tabela.ref)
            indice = [coluna.name for coluna in tabela.tableColumns].index("REALIZADO TOTAL")
            tabela.tableColumns[indice].name = "CÓDIGO DO PROJETO EMBRAPII2"
            ws.cell(inicio, primeira + indice).value = "CÓDIGO DO PROJETO EMBRAPII2"
            with tempfile.TemporaryDirectory() as pasta:
                caminho = Path(pasta) / "cabecalho-incorreto.xlsx"
                wb.save(caminho)
                with self.assertRaisesMessage(services.IngestaoError, "REALIZADO"):
                    self.importar(caminho)
        finally:
            wb.close()
        self.assertEqual(list(ProjetoFinanceiro.objects.order_by("pk").values_list("nome", "realizado")), antes)
        self.assertEqual(Acompanhamento.objects.count(), 1)
