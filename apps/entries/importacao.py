"""Importação de valores operacionais, limitada ao período e pilar autorizados."""
import csv
import io
from decimal import Decimal
from openpyxl import load_workbook
from django.core.exceptions import ValidationError
from django.db import transaction
from apps.audit.services import registrar_auditoria
from .services import salvar_ou_enviar

@transaction.atomic
def importar_operacionais(arquivo, periodo, pilar, usuario):
    if arquivo.size > 5 * 1024 * 1024:
        raise ValidationError("A planilha deve ter até 5 MB.")
    nome = arquivo.name.lower()
    if nome.endswith(".xlsx"):
        try:
            workbook = load_workbook(arquivo, read_only=True, data_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
        except Exception as exc:
            raise ValidationError("Planilha XLSX inválida.") from exc
    elif nome.endswith(".csv"):
        try:
            rows = list(csv.reader(io.StringIO(arquivo.read().decode("utf-8-sig")), delimiter=";"))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValidationError("CSV inválido. Use UTF-8 e separador ponto e vírgula.") from exc
    else:
        raise ValidationError("Envie um arquivo .xlsx ou .csv.")
    if not rows or [str(c or "").strip().lower() for c in rows[0][:3]] != ["codigo", "valor", "comentario"]:
        raise ValidationError("Cabeçalho esperado: codigo;valor;comentario.")
    indicadores = {i.codigo: i for i in pilar.indicadores.filter(ativo=True)}
    dados = {}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        codigo = str(row[0]).strip()
        ind = indicadores.get(codigo)
        if ind is None or str(ind.pk) in dados:
            raise ValidationError(f"Código desconhecido ou repetido no pilar: {codigo}.")
        valor = row[1] if len(row) > 1 else ""
        if isinstance(valor, (int, float, Decimal)):
            valor = str(valor).replace(".", ",")
        dados[str(ind.pk)] = {"valor": str(valor if valor is not None else ""), "comentario": str(row[2] or "") if len(row) > 2 else ""}
    total = salvar_ou_enviar(periodo, pilar, usuario, dados)
    registrar_auditoria(usuario, "IMPORTAR_LANCAMENTOS", "Periodo", periodo.pk, "planilha", "", arquivo.name, f"{total} rascunhos importados no pilar {pilar.codigo}.")
    return total
